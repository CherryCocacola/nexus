"""
DocumentExport 도구의 "포맷별 렌더러" 모듈.

한 줄 요약:
  모델(LLM)이 만들어낸 마크다운 유사 텍스트를 받아서, 실제로 열어볼 수 있는
  문서 파일(md / txt / docx / pptx / hwpx)로 변환·저장하는 함수들의 모음이다.

왜 이 파일이 존재하나:
  - DocumentExport 도구는 "무엇을 저장할지(내용)"와 "어떤 포맷으로 저장할지"를
    분리한다. 이 파일은 후자, 즉 "포맷별로 파일을 실제로 써 내려가는" 부분만 담당한다.
  - 지원 포맷은 전부 순수 파이썬 라이브러리(python-docx / python-pptx / python-hwpx)
    로만 만든다. 외부 서버·바이너리에 의존하지 않으므로 에어갭(폐쇄망)에서도 안전하다.

전체 설계(핵심 2가지):
  1) parse_blocks() 로 입력 텍스트를 공통 "블록" 목록으로 딱 한 번만 파싱한다.
     블록은 (제목 / 글머리표 / 문단) 3종뿐이라 단순하다. 각 포맷 렌더러는 이
     블록 목록을 "소비만" 하므로, 포맷마다 파싱 로직을 중복 구현하지 않아도 된다.
  2) 무거운 라이브러리(python-docx / pptx / hwpx) import 는 각 렌더러 함수 "안"에서
     지연(lazy) import 한다. 도구가 로드되는 시점에 그 라이브러리가 설치돼 있지 않아도
     프로세스가 죽지 않고, 실제로 해당 포맷을 요청받는 순간에만 import 가 일어난다.

입력 문법(모델이 자연스럽게 쓰는 마크다운의 부분집합):
  - "# / ## / ###"   → 제목(레벨 1~3, '#' 개수로 결정)
  - "- 항목" / "* 항목" → 글머리표(연속된 줄은 하나의 묶음으로 합쳐짐)
  - "**굵게**"        → 굵은 글씨(포맷이 지원하면 반영, 아니면 평문으로 처리)
  - 그 외 비어있지 않은 줄 → 일반 문단

외부에 노출하는 것(다른 모듈이 import 해서 쓰는 것):
  - RENDERERS   : 포맷명 → 렌더 함수 매핑(도구·다운로드 라우트가 공유)
  - MEDIA_TYPES : 포맷명 → HTTP MIME 타입 매핑(웹 다운로드 응답 헤더에 사용)

작성자: 이현수 / 작성일: 2026-07-05
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

# 블록(Block)의 표현 규칙 — parse_blocks() 가 만들고 각 렌더러가 읽는 공통 자료구조.
# 파이썬 튜플로 표현하며, 첫 원소(문자열)가 종류를 나타내는 태그다:
#   ("h", level:int, text:str)  — 제목. level 은 1~3.
#   ("bullet", items:list[str]) — 글머리표 한 묶음. 항목들의 리스트를 담는다.
#   ("p", text:str)             — 일반 문단.
# 별도 클래스를 두지 않고 튜플로 가볍게 표현한다(파싱→렌더 사이 단순 전달용이라 충분).
Block = tuple


def strip_md(text: str) -> str:
    """
    '**굵게**' 같은 마크다운 굵기 표시를 벗겨내 순수 텍스트만 남긴다.

    쓰이는 곳: 굵기를 표현할 수 없는 포맷(txt) 이나, 굵기를 지원하지 않는
    렌더 경로에서 '**' 기호가 그대로 노출되지 않도록 미리 제거하는 용도다.
    동작: 정규식으로 '**...**' 를 찾아 안쪽 내용(캡처 그룹 1)만 남긴다.
    """
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def parse_blocks(content: str) -> list[Block]:
    """
    마크다운 유사 텍스트를 공통 블록 목록(list[Block])으로 파싱한다.

    이 함수가 모든 포맷 렌더러의 "공통 입구"다. 한 번 파싱해 둔 블록 목록을
    docx / pptx / hwpx 렌더러가 각자 소비하므로 파싱을 중복하지 않는다.

    핵심 규칙:
      - 빈 줄을 만나면 지금까지 모아둔 글머리표 묶음을 확정(flush)한다.
      - 연속된 글머리표 줄("- ", "* ")은 하나의 ("bullet", [...]) 로 묶는다.
        각 포맷에서 <ul>/불릿 문단 그룹처럼 자연스럽게 렌더하기 위함이다.

    매개변수:
      content — 모델이 만든 원본 텍스트(여러 줄).
    반환:
      Block 튜플들의 리스트(위에서 정의한 h / bullet / p 세 종류).
    """
    blocks: list[Block] = []
    bullet_buf: list[str] = []  # 아직 확정되지 않은 글머리표 항목들을 잠시 담아두는 버퍼

    def flush_bullets() -> None:
        # 버퍼에 모아둔 글머리표 항목이 있으면, 하나의 bullet 블록으로 확정하고 버퍼를 비운다.
        # nonlocal 을 쓰는 이유: 바깥 지역변수 bullet_buf 를 이 안쪽 함수에서 재바인딩하기 때문.
        nonlocal bullet_buf
        if bullet_buf:
            blocks.append(("bullet", bullet_buf))
            bullet_buf = []

    # 입력을 한 줄씩 순회하며 블록으로 분류한다.
    for raw in content.splitlines():
        line = raw.rstrip()  # 줄 끝 공백 제거
        stripped = line.strip()  # 앞뒤 공백까지 제거한 판정용 문자열
        if not stripped:
            # 빈 줄 = 문단 경계. 모아둔 글머리표를 끊어주고 다음 줄로.
            flush_bullets()
            continue

        # 제목 판정: 줄 맨 앞의 '#' 1~6개 뒤에 공백이 오면 제목으로 본다.
        # 레벨은 '#' 개수지만 최대 3으로 제한한다(그보다 깊은 제목은 3으로 취급).
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush_bullets()
            level = min(len(m.group(1)), 3)
            blocks.append(("h", level, m.group(2).strip()))
            continue

        # 글머리표 판정: "- " 또는 "* " 로 시작하면 항목으로 보고 버퍼에 쌓는다.
        # (여기서 flush 하지 않는 이유: 연속된 항목을 한 묶음으로 합치기 위해서다.)
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            bullet_buf.append(m.group(1).strip())
            continue

        # 위 어디에도 해당하지 않으면 일반 문단.
        # 문단 앞에서 글머리표 묶음을 끊어준 뒤 문단 블록을 추가한다.
        flush_bullets()
        blocks.append(("p", stripped))

    # 입력 끝에 도달했을 때 아직 확정 안 된 글머리표가 남아 있으면 마지막으로 확정.
    flush_bullets()
    return blocks


# ─────────────────────────────────────────────
# 평문 계열: md / txt
#   외부 라이브러리 없이 그냥 텍스트 파일로 쓰는 가장 단순한 두 렌더러.
# ─────────────────────────────────────────────
def render_md(content: str, path: Path, title: str = "") -> None:
    """
    md 렌더러 — 입력 자체가 이미 마크다운이므로 별도 변환 없이 그대로 저장한다.

    title 이 주어지면 문서 맨 앞에 "# 제목" 한 줄을 얹어준다.
    매개변수: content(본문), path(저장 경로), title(선택 제목). 반환: 없음(파일로 저장).
    """
    body = f"# {title}\n\n{content}" if title else content
    path.write_text(body, encoding="utf-8")


def render_txt(content: str, path: Path, title: str = "") -> None:
    """
    txt 렌더러 — 굵기 표시(**)를 벗겨낸 순수 평문으로 저장한다.

    txt 는 서식을 표현할 수 없으므로, strip_md() 로 마크다운 굵기 기호를 지운 뒤 쓴다.
    title 이 있으면 맨 앞에 제목 줄을 붙인다. 반환: 없음(파일로 저장).
    """
    body = f"{title}\n\n{content}" if title else content
    path.write_text(strip_md(body), encoding="utf-8")


# ─────────────────────────────────────────────
# DOCX (python-docx)
#   워드 문서(.docx) 생성. 제목/글머리표/굵기 서식까지 반영한다.
# ─────────────────────────────────────────────
def _docx_add_runs(paragraph: Any, text: str) -> None:
    """
    한 줄의 텍스트를 '**굵게**' 기준으로 쪼개 워드 문단(run)에 이어 붙인다.

    python-docx 에서 서식(굵기)은 문단(paragraph) 안의 "run" 단위로 지정한다.
    그래서 굵은 구간과 일반 구간을 각각 별도 run 으로 나눠 추가해야 한다.

    동작 원리:
      re.split 로 '**...**' 를 기준 삼아 나누면, 캡처된 굵기 대상은 항상
      홀수 인덱스(1, 3, 5...) 에 오고 나머지는 일반 텍스트가 된다.
    매개변수: paragraph(대상 문단 객체), text(원본 한 줄). 반환: 없음.
    """
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for i, part in enumerate(parts):
        if part == "":
            continue  # 굵기 기호가 줄 양끝에 있으면 빈 조각이 생기므로 건너뛴다.
        run = paragraph.add_run(part)
        if i % 2 == 1:  # 캡처된 굵기 대상은 홀수 인덱스 → 굵게 처리
            run.bold = True


def render_docx(content: str, path: Path, title: str = "") -> None:
    """
    DOCX(.docx) 생성 렌더러 — 제목/글머리표/굵기 서식을 반영한다.

    흐름:
      1) 문서 객체 생성, title 이 있으면 최상위 제목(level 0)으로 추가.
      2) parse_blocks() 결과를 순회하며 블록 종류별로 워드 요소를 추가:
         - h      → add_heading(레벨 그대로)
         - bullet → 항목마다 "List Bullet" 스타일 문단
         - p      → 일반 문단(굵기 서식은 _docx_add_runs 로 처리)
      3) 파일로 저장.
    반환: 없음(path 에 .docx 파일 생성).
    """
    from docx import Document  # lazy import — 실제 docx 요청 시에만 로드

    doc = Document()
    if title:
        doc.add_heading(title, level=0)

    for block in parse_blocks(content):
        if block[0] == "h":
            doc.add_heading(block[2], level=block[1])
        elif block[0] == "bullet":
            for item in block[1]:
                p = doc.add_paragraph(style="List Bullet")
                _docx_add_runs(p, item)
        else:  # "p" — 일반 문단
            p = doc.add_paragraph()
            _docx_add_runs(p, block[1])

    doc.save(str(path))


# ─────────────────────────────────────────────
# PPTX (python-pptx)
#   파워포인트(.pptx) 생성. 제목 블록을 "슬라이드 나누는 기준"으로 삼는다.
# ─────────────────────────────────────────────
def render_pptx(content: str, path: Path, title: str = "") -> None:
    """
    PPTX(.pptx) 생성 렌더러.

    슬라이드 구성 규칙:
      - 맨 앞에 제목 슬라이드 1장을 둔다(제목 = title, 없으면 첫 제목 블록 텍스트).
      - 이후 제목 블록(#/##/###)을 만날 때마다 새 본문 슬라이드를 시작하고,
        그 뒤에 이어지는 글머리표/문단을 그 슬라이드의 본문으로 채운다.
      - 제목 블록이 하나도 없으면 본문 슬라이드 1장에 전부 담는다.

    구현 메모: 파워포인트 본문은 text_frame 안의 문단들로 채운다. 첫 줄은 이미
    존재하는 기본 문단을 재사용하고, 둘째 줄부터 add_paragraph 로 추가한다.
    반환: 없음(path 에 .pptx 파일 생성).
    """
    from pptx import Presentation  # lazy import — 실제 pptx 요청 시에만 로드

    prs = Presentation()
    blocks = parse_blocks(content)

    # 제목 슬라이드 — 레이아웃 0(제목 슬라이드).
    # 덱 제목은 title 우선, 없으면 첫 제목 블록, 그것도 없으면 "문서".
    first_heading = next((b[2] for b in blocks if b[0] == "h"), "")
    deck_title = title or first_heading or "문서"
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = deck_title
    if len(s0.placeholders) > 1:
        # 부제목 자리(placeholder 1)가 있으면 생성 출처 문구를 넣는다.
        s0.placeholders[1].text = "Nexus 생성 문서"

    # 본문 슬라이드들
    body_tf = None  # 지금 채우고 있는 본문 슬라이드의 text_frame(아직 없으면 None)

    def new_body_slide(heading_text: str):
        # 레이아웃 1(제목+본문)로 새 슬라이드를 만들고, 본문 text_frame 을 비워 반환한다.
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading_text
        tf = slide.placeholders[1].text_frame
        tf.clear()
        return tf

    def add_body_line(tf, text: str, first_flag: list[bool]) -> None:
        # 본문 한 줄을 추가한다. text_frame 의 첫 줄은 기본 문단을 재사용하고,
        # 그 이후 줄부터는 add_paragraph 로 새 문단을 만든다.
        # first_flag 를 리스트로 넘기는 이유: 안쪽에서 값을 바꿔 바깥과 공유하기 위해(참조 전달).
        clean = strip_md(text)  # 슬라이드 본문엔 굵기 기호를 노출하지 않는다.
        if first_flag[0]:
            tf.paragraphs[0].text = clean
            first_flag[0] = False
        else:
            tf.add_paragraph().text = clean

    first_line = [True]  # 현재 슬라이드에서 아직 첫 줄을 안 썼는지 표시하는 플래그
    for block in blocks:
        if block[0] == "h":
            # 제목 블록 = 새 슬라이드 시작 신호.
            body_tf = new_body_slide(block[2])
            first_line = [True]
        else:
            if body_tf is None:
                # 제목 블록 없이 본문(문단/글머리표)이 먼저 온 경우 → 덱 제목으로 슬라이드 개설.
                body_tf = new_body_slide(deck_title)
                first_line = [True]
            if block[0] == "bullet":
                for item in block[1]:
                    add_body_line(body_tf, item, first_line)
            else:  # "p" — 일반 문단
                add_body_line(body_tf, block[1], first_line)

    prs.save(str(path))


# ─────────────────────────────────────────────
# HWPX (python-hwpx)
#   한글(.hwpx) 생성. hwpx.builder 의 고수준 API 로 문서 트리를 조립한다.
# ─────────────────────────────────────────────
def render_hwpx(content: str, path: Path, title: str = "") -> None:
    """
    HWPX(한글, .hwpx) 생성 렌더러 — hwpx.builder 고수준 API 를 사용한다.

    흐름:
      1) 자식 요소 리스트(children)를 준비하고, title 이 있으면 Heading(레벨 1)으로 시작.
      2) parse_blocks() 결과를 순회하며 블록을 hwpx.builder 요소로 변환:
         - h      → Heading(레벨 그대로)
         - bullet → Bullet(항목 리스트)
         - p      → Paragraph
         (hwpx 경로는 굵기 서식을 별도 처리하지 않으므로 strip_md 로 기호를 제거해 넣는다.)
      3) children 을 하나의 Section 에 담아 Document 를 만들고 파일로 저장.
    반환: 없음(path 에 .hwpx 파일 생성).
    """
    from hwpx.builder import Bullet, Document, Heading, Paragraph, Section  # lazy

    children: list[Any] = []
    if title:
        children.append(Heading(level=1, text=title))

    for block in parse_blocks(content):
        if block[0] == "h":
            children.append(Heading(level=block[1], text=strip_md(block[2])))
        elif block[0] == "bullet":
            children.append(Bullet(items=[strip_md(x) for x in block[1]]))
        else:  # "p" — 일반 문단
            children.append(Paragraph(text=strip_md(block[1])))

    doc = Document(sections=[Section(children=children)])
    doc.save_to_path(str(path))


# 포맷명 → 렌더 함수 매핑.
# 도구 실행부와 웹 다운로드 라우트가 "어떤 포맷을 어떤 함수로 처리할지"를
# 여기 한 곳에서만 참조하도록 하는 단일 진실원(single source of truth)이다.
RENDERERS: dict[str, Any] = {
    "md": render_md,
    "txt": render_txt,
    "docx": render_docx,
    "pptx": render_pptx,
    "hwpx": render_hwpx,
}

# 포맷명 → HTTP MIME(미디어) 타입 매핑.
# 웹에서 파일을 내려줄 때 Content-Type 헤더 값으로 사용한다.
MEDIA_TYPES: dict[str, str] = {
    "md": "text/markdown; charset=utf-8",
    "txt": "text/plain; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "hwpx": "application/hwp+zip",
    # ImageGenerate 도구가 저장한 PNG를 웹 다운로드 라우트가 image/png 로 서빙해
    # <img> 미리보기가 올바르게 표시되도록 한다(없으면 octet-stream 폴백).
    "png": "image/png",
}
