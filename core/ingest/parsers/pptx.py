"""
PPTX 파서 — python-pptx(MIT)로 PowerPoint 파일을 "구조 트리"로 변환하는 모듈.

이 파일이 하는 일 (한 문장):
  .pptx 파일 하나를 받아, 슬라이드/도형/표/그림/문단을 의미 단위 노드로 쪼갠
  DocumentTree 를 만들어 돌려준다. 이 트리는 뒤이어 RAG 임베딩·검색 파이프라인의
  입력이 된다(이 파일 자체는 임베딩·검색을 하지 않는다).

왜 PPTX 가 1차 구현 대상인가 (v7.3 Part 3.1):
  PPTX 는 슬라이드/도형/텍스트박스/표/그림을 "네이티브 객체"로 제공한다.
  즉 레이아웃 모델이나 OCR 없이도 구조를 그대로 읽어올 수 있어 가장 쉽다.
  "충돌 없는 임베딩"의 핵심인 읽기 순서 복원도, 각 도형의 (top,left) 좌표를
  정렬하는 것만으로 해결된다(좌측 텍스트박스와 우측 텍스트박스가 뒤섞이지
  않게 함).

처리 개요:
  파일 → 슬라이드 순회 → 슬라이드마다 SLIDE 노드 생성
    └ 슬라이드 안 도형을 (top,left) 정렬 → 읽기 순서대로 자식 노드 생성
        - 텍스트박스/플레이스홀더 → TEXTBOX 또는 PARAGRAPH/LIST_ITEM
        - 표 → TABLE (행/열 보존)
        - 그림 → FIGURE_CAPTION (대체텍스트/캡션이 있으면)
        - 그룹(GroupShape) → 재귀 순회

핵심 구성 요소:
  - PptxParser: 이 모듈의 유일한 공개 클래스. DocumentParser(ABC) 를 상속한다.
      · can_parse(path): 이 파서가 해당 파일을 처리할 수 있는지 사전 판정.
      · parse(path):     실제 변환. 비동기(async) 진입점이다.
  - 나머지 _언더스코어 메서드들은 모두 parse 를 돕는 내부 헬퍼다(외부 노출 X).

fail-soft (anti-pattern #8):
  도형 하나를 읽다 실패해도 전체를 중단하지 않는다. 실패 사유를 warnings 에
  모으고 나머지를 계속 처리한다. bare except 는 쓰지 않고 구체 예외만 잡는다.
  → 손상된 슬라이드가 한 장 있어도 나머지 슬라이드는 정상적으로 트리에 담긴다.

의존성 방향 (P2): core.ingest.types / parser_base 만 의존. core/rag·model 무관.
에어갭: python-pptx 는 로컬 파일만 읽는다(외부 네트워크 없음).

작성자: 이현수 / 작성일: 2026-07-05
"""

from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

# 이 파서가 상속할 추상 베이스(DocumentParser)와, 만들어 낼 데이터 모델들.
# - DocumentParser: 모든 포맷 파서가 구현해야 하는 공통 인터페이스(ABC).
# - DocumentTree/DocumentNode: 파싱 결과를 담는 트리 구조(Pydantic 모델).
# - ElementType: 노드의 종류(SLIDE/TABLE/PARAGRAPH/LIST_ITEM/FIGURE_CAPTION 등) enum.
from core.ingest.parser_base import DocumentParser
from core.ingest.types import DocumentNode, DocumentTree, ElementType

# 모듈 전용 로거. 규칙(P7)에 따라 "nexus.{module}" 네임스페이스를 사용한다.
logger = logging.getLogger("nexus.ingest.parsers.pptx")

# PPTX(OOXML)는 ZIP 컨테이너다. 매직바이트는 "PK\x03\x04"(ZIP 시그니처).
# can_parse() 에서 확장자뿐 아니라 이 시그니처도 확인해 fail-closed 한다.
# (확장자만 .pptx 로 바꾼 가짜/손상 파일을 미리 걸러내기 위함.)
_ZIP_MAGIC = b"PK\x03\x04"

# 좌표가 없는 도형(placeholder 등)을 정렬에서 맨 뒤로 보내기 위한 큰 값(≈10^18).
# 정렬 키에서 top/left 가 None 인 도형에 이 값을 대입하면 항상 맨 뒤로 밀린다.
# v7.3 Part 3.3 의 _ordered_shapes 발췌와 동일한 의도.
_NO_COORD = 10**18


class PptxParser(DocumentParser):
    """
    PowerPoint(.pptx) 파일을 구조 트리(DocumentTree)로 파싱하는 파서.

    DocumentParser(ABC) 를 상속하며, 인제스트 파이프라인은 파일 확장자를 보고
    등록된 여러 파서 중 이 클래스를 골라 parse() 를 호출한다.

    네이티브 구조를 그대로 읽으므로 GPU 가 필요 없다(requires_gpu=False).
    스캔 PDF 처럼 OCR/레이아웃 모델이 필요한 포맷과 달리, PPTX 는 도형·텍스트가
    이미 객체로 존재하므로 좌표 정렬만으로 읽기 순서를 복원할 수 있다.

    주요 진입점:
      - supported_extensions / requires_gpu: 파서 등록·선택용 메타 정보(property).
      - can_parse(path):  처리 가능 여부 사전 판정(확장자 + 매직바이트).
      - parse(path):      실제 변환. async — 파이프라인이 await 로 호출한다.
    """

    # ─── 정체성/플래그 ───

    @property
    def supported_extensions(self) -> tuple[str, ...]:
        # 이 파서가 처리하는 확장자 목록. 파이프라인이 파서를 고를 때 참조한다.
        # .pptx 만 처리. .ppt(구 바이너리 포맷)는 python-pptx 미지원이라 대상 아님.
        return (".pptx",)

    @property
    def requires_gpu(self) -> bool:
        # 네이티브 구조 파싱이라 GPU 레이아웃/OCR 모델이 전혀 필요 없다.
        # 파이프라인은 이 값으로 GPU 자원 배정 여부를 결정한다(False → CPU 만).
        return False

    # ─── 처리 가능 여부 (fail-closed) ───

    def can_parse(self, path: Path) -> bool:
        """
        이 파일을 PptxParser 로 처리할 수 있는지 사전 판정한다(실제 파싱 전 필터).

        판정 기준(둘 다 만족해야 True):
          1. 확장자가 .pptx 인가 (대소문자 무시).
          2. 파일 앞 4바이트가 ZIP 시그니처(_ZIP_MAGIC)인가 — PPTX 는 ZIP 컨테이너.

        매개변수:
          path: 검사할 파일 경로.
        반환:
          bool — 처리 가능하면 True.

        fail-closed: 파일을 읽다 문제가 생기거나 시그니처가 다르면 False 로 막는다.
        (의심스러우면 "처리 가능"이 아니라 "처리 불가"로 기우는 안전한 기본값.)
        """
        # 1) 확장자 필터 — 가장 싼 검사부터 먼저 수행해 불필요한 파일 I/O 를 피한다.
        if path.suffix.lower() != ".pptx":
            return False
        # 2) 매직바이트 검사 — 앞 4바이트만 읽어 ZIP 컨테이너인지 확인.
        try:
            with path.open("rb") as f:
                head = f.read(4)
        except OSError as e:
            # 파일을 열 수 없음(권한/부재/잠김 등) — 처리 불가로 간주(fail-closed).
            logger.debug("can_parse 파일 열기 실패: %s (%s)", path, e)
            return False
        return head == _ZIP_MAGIC

    # ─── 핵심: parse ───

    async def parse(self, path: Path) -> DocumentTree:
        """
        .pptx 를 슬라이드 단위 구조 트리(DocumentTree)로 변환하는 메인 진입점.

        전체 흐름(3단계):
          1) Presentation 로드 — python-pptx 가 ZIP 을 풀어 슬라이드 객체를 준다.
          2) 슬라이드를 순회하며 각 슬라이드를 SLIDE 노드(+자식들)로 변환.
          3) 문서 제목을 정하고(첫 슬라이드 제목 → 없으면 파일명) 트리를 조립.

        매개변수:
          path: 파싱할 .pptx 파일 경로.
        반환:
          DocumentTree(title, source_path, nodes=SLIDE 노드들, doc_format="pptx",
          warnings=fail-soft 경고 목록).

        비동기(async)인 이유: 인제스트 파이프라인이 여러 파일을 동시에 처리할 수
        있도록 공통 인터페이스가 async 로 정의돼 있다(이 구현 자체는 CPU 작업).

        치명적 실패(파일 손상 등)는 예외를 던지지 않고 빈 트리 + 경고로 표현한다
        (한 파일이 망가져도 전체 인제스트 배치가 멈추지 않도록 — 파이프라인 보호).
        """
        # fail-soft 로 모은 경고 메시지들. 최종적으로 트리의 warnings 에 담긴다.
        warnings: list[str] = []

        # 1) 파일 로드 — 손상/형식 오류는 부분 결과(빈 트리)+경고로 처리.
        try:
            presentation = Presentation(str(path))
        except (PackageNotFoundError, OSError, ValueError, KeyError) as e:
            # python-pptx 는 손상/비정상 파일에서 PackageNotFoundError 를 던진다.
            # 이 예외는 PythonPptxError→Exception 계열이라 OSError/ValueError/KeyError
            # 어디에도 속하지 않으므로, 반드시 명시적으로 포착해야 fail-soft(빈 트리+
            # 경고)가 보장된다. bare except 금지 — 구체 예외만 포착.
            msg = f"PPTX 로드 실패: {type(e).__name__}: {e}"
            logger.warning("%s (%s)", msg, path)
            return DocumentTree(
                title=path.stem,
                source_path=str(path),
                nodes=(),
                doc_format="pptx",
                warnings=(msg,),
            )

        # 2) 슬라이드 순회 → SLIDE 노드 생성.
        #    슬라이드 번호는 사람이 세는 1부터 시작(enumerate start=1).
        slide_nodes: list[DocumentNode] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            slide_node = self._parse_slide(slide, slide_index, warnings)
            slide_nodes.append(slide_node)

        # 3) 문서 제목: 첫 슬라이드 제목 placeholder → 없으면 파일명(stem).
        #    (stem = 확장자 뺀 파일명. 예: "보안교육.pptx" → "보안교육")
        title = self._document_title(slide_nodes) or path.stem

        return DocumentTree(
            title=title,
            source_path=str(path),
            nodes=tuple(slide_nodes),
            doc_format="pptx",
            warnings=tuple(warnings),
        )

    # ─────────────────────────────────────────
    # 내부 헬퍼
    # ─────────────────────────────────────────

    def _parse_slide(self, slide, slide_index: int, warnings: list[str]) -> DocumentNode:
        """
        슬라이드 1장을 SLIDE 노드(+자식 노드들)로 변환한다.

        매개변수:
          slide:       python-pptx 의 Slide 객체.
          slide_index: 1부터 시작하는 슬라이드 번호(페이지 번호로도 사용).
          warnings:    fail-soft 경고를 누적할 리스트(호출부와 공유, 계속 append).
        반환:
          DocumentNode(element_type=SLIDE) — children 에 슬라이드 내용 노드들이 담김.

        단계:
          1. 슬라이드 제목 placeholder 에서 heading(슬라이드 제목)을 추출.
             이 제목이 슬라이드 내 모든 자식 노드의 heading_path 가 된다
             (문맥 연결 — 같은 슬라이드 내용은 같은 제목 맥락을 공유).
          2. 도형을 (top,left) 좌표로 정렬해 읽기 순서를 복원.
          3. 각 도형을 종류별 노드로 변환(그룹은 재귀).
        """
        # 1) 슬라이드 제목 — heading_path 의 한 단계가 된다.
        slide_title = self._slide_title(slide)
        # heading_path 는 "슬라이드 N" 또는 실제 제목을 사용한다.
        # 제목이 있으면 그것을, 없으면 "슬라이드 N" 라벨로 맥락을 표시.
        heading = slide_title or f"슬라이드 {slide_index}"
        heading_path: tuple[str, ...] = (heading,)

        # 2)+3) 읽기 순서대로 도형을 자식 노드로 변환.
        children = self._shapes_to_nodes(
            shapes=self._ordered_shapes(slide.shapes),
            heading_path=heading_path,
            page=slide_index,
            warnings=warnings,
            skip_title=slide_title,  # 제목 placeholder 텍스트 중복 방지
        )

        # SLIDE 노드 조립.
        #   text:  제목이 있으면 그 제목을, 없으면 빈 문자열(_document_title 이 참조).
        #   heading_path: SLIDE 는 경로의 루트라 자기 heading 을 넣지 않는다(빈 튜플).
        #   order: 슬라이드 순서 = 슬라이드 번호. children 은 읽기 순서로 정렬된 노드들.
        return DocumentNode(
            element_type=ElementType.SLIDE,
            text=heading if slide_title else "",  # 제목이 있으면 SLIDE 노드 텍스트로
            heading_path=(),  # SLIDE 자신은 경로의 루트
            page=slide_index,
            order=slide_index,
            children=tuple(children),
        )

    @staticmethod
    def _ordered_shapes(shapes) -> list:
        """
        도형을 (top, left) 좌표순으로 정렬해 읽기 순서를 복원한다.

        왜 필요한가: PPTX 안 도형의 저장 순서는 "화면상 읽는 순서"와 무관하다.
        위→아래, 같은 높이면 좌→우 로 정렬해야 사람이 읽는 순서가 재현된다.
        (staticmethod 인 이유: self 상태에 의존하지 않는 순수 정렬 유틸이라서.)

        매개변수:
          shapes: python-pptx 의 도형 컬렉션(슬라이드의 shapes 또는 그룹의 shapes).
        반환:
          (top, left) 오름차순으로 정렬된 도형 리스트.

        v7.3 Part 3.3 발췌와 동일한 의도:
          좌표가 없는 도형(레이아웃 placeholder 등)은 매우 큰 값(_NO_COORD)으로
          처리해 맨 뒤로 보낸다. 이렇게 하면 좌측 텍스트박스와 우측 텍스트박스가
          읽기 순서 없이 뒤섞이는 "충돌"을 막는다.

        sorted 는 안정 정렬이므로 좌표가 같은 도형은 원래 순서를 유지한다.
        """

        def sort_key(shape):
            # getattr 로 안전 접근 — 일부 도형 타입은 top/left 가 None 일 수 있다.
            # None 이면 _NO_COORD 로 치환해 정렬 시 맨 뒤로 밀리게 한다.
            top = getattr(shape, "top", None)
            left = getattr(shape, "left", None)
            return (
                top if top is not None else _NO_COORD,
                left if left is not None else _NO_COORD,
            )

        return sorted(shapes, key=sort_key)

    def _shapes_to_nodes(
        self,
        shapes: list,
        heading_path: tuple[str, ...],
        page: int,
        warnings: list[str],
        skip_title: str | None,
    ) -> list[DocumentNode]:
        """
        정렬된 도형 목록을 자식 노드 목록으로 변환한다(읽기 순서 유지).

        이 메서드가 파서의 "심장"이다: 도형 종류를 판별해 알맞은 노드로 분기한다.
        그룹 도형을 만나면 자기 자신을 재귀 호출해 내부 도형까지 펼친다.

        매개변수:
          shapes:       이미 _ordered_shapes 로 정렬된 도형 리스트.
          heading_path: 이 도형들이 속한 제목 경로(슬라이드 제목 맥락).
          page:         슬라이드 번호(경고 메시지·노드 page 필드에 사용).
          warnings:     fail-soft 경고 누적 리스트.
          skip_title:   제목 placeholder 텍스트. 이와 같은 문단은 중복이라 건너뜀.
        반환:
          DocumentNode 리스트(읽기 순서대로 order 0,1,2... 부여됨).

        도형 종류별 매핑:
          - GROUP        → 재귀 순회(내부 도형도 좌표 정렬)
          - TABLE        → TABLE 노드 (행/열 보존)
          - PICTURE      → FIGURE_CAPTION (대체텍스트가 있을 때만)
          - 그 외 텍스트  → TEXTBOX/PARAGRAPH/LIST_ITEM

        order 는 정렬된 순서대로 0,1,2... 를 부여해 읽기 순서를 보존한다.
        """
        nodes: list[DocumentNode] = []
        order = 0  # 자식 노드에 매길 읽기 순서 번호. 노드를 추가할 때마다 +1.

        for shape in shapes:
            # 도형 타입 판별 — shape_type 접근 자체가 실패할 수 있어 보호한다.
            # (일부 비정상 도형은 shape_type 조회에서 예외를 던진다 → 건너뜀.)
            try:
                shape_type = shape.shape_type
            except (AttributeError, ValueError) as e:
                warnings.append(
                    f"슬라이드 {page}: 도형 타입 판별 실패 — 건너뜀 ({type(e).__name__}: {e})"
                )
                continue

            # (a) 그룹 도형 — 내부 도형을 재귀적으로 좌표 정렬해 순회.
            #     그룹은 여러 도형을 묶은 컨테이너라, 내부를 펼쳐 평평하게 만든다.
            if shape_type == MSO_SHAPE_TYPE.GROUP:
                try:
                    inner = self._ordered_shapes(shape.shapes)
                except (AttributeError, ValueError) as e:
                    warnings.append(
                        f"슬라이드 {page}: 그룹 내부 순회 실패 — 건너뜀 ({type(e).__name__}: {e})"
                    )
                    continue
                # 재귀 호출: 그룹 내부는 자체 좌표계라 skip_title 은 넘기지 않는다.
                child_nodes = self._shapes_to_nodes(
                    inner, heading_path, page, warnings, skip_title=None
                )
                # 재귀가 매긴 내부 order 는 버리고, 부모 흐름의 연속 번호로 재부여한다.
                # DocumentNode 는 불변(frozen)이라 model_copy 로 새 노드를 만든다.
                for cn in child_nodes:
                    nodes.append(cn.model_copy(update={"order": order}))
                    order += 1
                continue

            # (b) 표 도형 — 행/열을 보존한 TABLE 노드 1개.
            #     파싱 실패 시 _table_node 가 None 을 주므로 그때는 노드를 안 넣는다.
            if shape_type == MSO_SHAPE_TYPE.TABLE:
                table_node = self._table_node(shape, heading_path, page, order, warnings)
                if table_node is not None:
                    nodes.append(table_node)
                    order += 1
                continue

            # (c) 그림 도형 — 대체텍스트/캡션이 있으면 FIGURE_CAPTION.
            #     캡션이 없는 순수 장식 그림은 검색 가치가 없어 노드를 만들지 않는다.
            if shape_type == MSO_SHAPE_TYPE.PICTURE:
                caption = self._picture_caption(shape)
                if caption:
                    nodes.append(
                        DocumentNode(
                            element_type=ElementType.FIGURE_CAPTION,
                            text=caption,
                            heading_path=heading_path,
                            page=page,
                            order=order,
                        )
                    )
                    order += 1
                continue

            # (d) 텍스트가 있는 그 외 도형(텍스트박스/플레이스홀더/오토셰이프).
            #     한 도형이 여러 문단을 가질 수 있어 노드 여러 개가 나올 수 있다.
            text_nodes = self._text_shape_nodes(shape, heading_path, page, order, skip_title)
            for tn in text_nodes:
                nodes.append(tn)
                order += 1

        return nodes

    def _table_node(
        self,
        shape,
        heading_path: tuple[str, ...],
        page: int,
        order: int,
        warnings: list[str],
    ) -> DocumentNode | None:
        """
        표 도형을 TABLE 노드로 변환한다(행/열 구조 보존).

        매개변수:
          shape:        MSO_SHAPE_TYPE.TABLE 인 도형.
          heading_path: 소속 제목 경로.  page: 슬라이드 번호.  order: 읽기 순서 번호.
          warnings:     fail-soft 경고 누적 리스트.
        반환:
          DocumentNode(TABLE) — 내용이 비었거나 파싱 실패면 None(호출부가 노드 생략).

        직렬화 형식: 각 행을 " | " 로 구분된 셀로, 행은 줄바꿈으로 연결한다.
        예) "이름 | 부서\n홍길동 | 보안팀". 이렇게 하면 행/열 의미가 보존되어
        검색·표시에 활용할 수 있다(v7.3 Part 1.1 — 표 셀이 본문과 섞이지 않게).

        표 접근 실패는 fail-soft: 경고를 남기고 None 반환(노드 생략).
        """
        try:
            table = shape.table
            rows_text: list[str] = []
            for row in table.rows:
                # 각 셀의 텍스트를 좌→우 순서로 모아 " | " 로 연결.
                cells = [cell.text.strip() for cell in row.cells]
                rows_text.append(" | ".join(cells))
        except (AttributeError, ValueError) as e:
            warnings.append(f"슬라이드 {page}: 표 파싱 실패 — 건너뜀 ({type(e).__name__}: {e})")
            return None

        # 완전히 빈 행은 제외하고, 남은 행을 줄바꿈으로 이어 하나의 텍스트로 만든다.
        content = "\n".join(rt for rt in rows_text if rt.strip())
        if not content:
            # 셀이 전부 비어 검색 가치가 없는 표 — 노드를 만들지 않는다.
            return None

        return DocumentNode(
            element_type=ElementType.TABLE,
            text=content,
            heading_path=heading_path,
            page=page,
            order=order,
        )

    @staticmethod
    def _picture_caption(shape) -> str:
        """
        그림 도형에서 대체텍스트(alt text) 또는 캡션을 추출한다.

        매개변수:
          shape: MSO_SHAPE_TYPE.PICTURE 인 도형.
        반환:
          캡션 문자열. 대체텍스트가 없으면 빈 문자열(호출부가 노드를 생략).

        python-pptx 는 도형 이름/대체텍스트를 노출한다. 대체텍스트가 가장
        의미 있는 캡션이므로 우선 사용하고, 없으면 빈 문자열(노드 생략).

        대체텍스트는 XML 의 cNvPr 요소 descr 속성에 들어있다. 안전 접근 후
        없으면 빈 문자열을 돌려준다.
        """
        # 공개 API 에 alt text 접근자가 없어 내부(_element) XML 노드로 직접 읽는다.
        # _nvXxPr.cNvPr 의 "descr" 속성이 PowerPoint 의 대체텍스트다.
        # 내부 속성 접근이라 도형 종류에 따라 실패할 수 있어 예외를 넓게 감싼다.
        try:
            descr = shape._element._nvXxPr.cNvPr.get("descr", "")
        except (AttributeError, KeyError, ValueError):
            descr = ""
        return (descr or "").strip()

    def _text_shape_nodes(
        self,
        shape,
        heading_path: tuple[str, ...],
        page: int,
        order: int,
        skip_title: str | None,
    ) -> list[DocumentNode]:
        """
        텍스트를 가진 도형(텍스트박스/플레이스홀더 등)을 노드 목록으로 변환한다.

        매개변수:
          shape:        텍스트프레임을 가질 수 있는 도형.
          heading_path: 소속 제목 경로.  page: 슬라이드 번호.
          order:        이 도형의 첫 문단에 부여할 시작 order(문단마다 +1).
          skip_title:   슬라이드 제목 텍스트. 같은 문단은 중복이라 제외.
        반환:
          문단별 DocumentNode 리스트(빈 문단·제목 중복은 제외됨).

        규칙:
          - 도형에 텍스트프레임이 없으면 빈 목록.
          - 제목 placeholder 와 동일한 텍스트(skip_title)는 중복이므로 제외
            (이미 슬라이드 제목/heading 으로 반영됨).
          - 문단(paragraph)별로 노드를 만들되, 글머리표(level>0)는 LIST_ITEM,
            그 외는 PARAGRAPH 로 분류한다.
          - 도형 전체는 TEXTBOX 의미를 가지나, 검색 정밀도를 위해 문단 단위로
            쪼개 노드를 만든다(빈 문단은 생략).

        텍스트 접근 실패는 빈 목록 반환(fail-soft — 호출부는 계속 진행).
        """
        # 텍스트프레임이 없는 도형(순수 그림·선 등)은 처리 대상이 아니다 → 빈 목록.
        try:
            if not shape.has_text_frame:
                return []
            text_frame = shape.text_frame
        except (AttributeError, ValueError):
            return []

        nodes: list[DocumentNode] = []
        local_order = order  # order 를 직접 안 건드리고 지역 변수로 증가시킨다.
        for para in text_frame.paragraphs:
            # 문단 안 run(서식 조각)들의 텍스트를 이어붙여 한 문단 텍스트로 만든다.
            try:
                para_text = "".join(run.text for run in para.runs).strip()
            except (AttributeError, ValueError):
                continue
            if not para_text:
                # 빈 문단(장식용 줄바꿈 등)은 건너뛴다.
                continue
            # 제목과 동일하면 중복 — 제외(이미 슬라이드 heading 으로 반영됨).
            if skip_title and para_text == skip_title.strip():
                continue

            # 들여쓰기 level 이 0보다 크면 글머리표(불릿) 항목으로 간주해 LIST_ITEM.
            # getattr 로 안전 접근하고, None 이면 0 으로 취급(or 0).
            level = getattr(para, "level", 0) or 0
            element_type = ElementType.LIST_ITEM if level > 0 else ElementType.PARAGRAPH

            nodes.append(
                DocumentNode(
                    element_type=element_type,
                    text=para_text,
                    heading_path=heading_path,
                    page=page,
                    order=local_order,
                )
            )
            local_order += 1

        return nodes

    @staticmethod
    def _slide_title(slide) -> str | None:
        """
        슬라이드 제목 placeholder 에서 제목 텍스트를 추출한다.

        매개변수:
          slide: python-pptx 의 Slide 객체.
        반환:
          제목 문자열, 또는 제목이 없으면 None.

        python-pptx 의 slide.shapes.title 은 제목 placeholder 가 있으면
        그 도형을, 없으면 None 을 반환한다. 안전하게 접근한다.
        (title 조회와 text 조회를 각각 try 로 감싸 어느 쪽이 실패해도 None 반환.)
        """
        try:
            title_shape = slide.shapes.title
        except (AttributeError, ValueError):
            return None
        if title_shape is None:
            return None
        try:
            text = (title_shape.text or "").strip()
        except (AttributeError, ValueError):
            return None
        # 공백만 있는 제목은 없는 것으로 취급 → None.
        return text or None

    @staticmethod
    def _document_title(slide_nodes: list[DocumentNode]) -> str | None:
        """
        문서 제목 후보를 첫 슬라이드의 텍스트(제목)에서 가져온다.

        매개변수:
          slide_nodes: parse() 가 만든 SLIDE 노드 리스트.
        반환:
          가장 앞에서 찾은 비어있지 않은 SLIDE 제목, 없으면 None.

        _parse_slide 에서 제목이 있으면 SLIDE 노드의 text 에 담아 두었으므로
        첫(=텍스트가 있는 첫) SLIDE 노드의 text 를 사용한다.
        없으면 None → parse() 가 파일명(stem)으로 폴백한다.
        """
        for node in slide_nodes:
            if node.element_type == ElementType.SLIDE and node.text.strip():
                return node.text.strip()
        return None
