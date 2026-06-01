"""
core.ingest.pipeline — DocumentIngestPipeline 통합 테스트 (v7.3 문서 인제스트).

전체 흐름(파싱→청킹→임베딩→적재)을 fake 의존성으로 e2e 검증한다:
  - fake model_provider: embed(texts) 만 구현(고정 차원 벡터 반환).
  - 인메모리 KnowledgeStore(pg_pool=None): DB 없이 UPSERT 동작 검증.

검증 의도:
  1. dry_run=True → 적재 없이 chunk_count 등 요약만 반환(부작용 0).
  2. dry_run=False → KnowledgeStore 에 KnowledgeEntry 적재(source="docingest"),
     metadata 에 element_type/heading_path/page 보존, embedding 은 tuple.
  3. 멱등성 — 같은 파일 2회 ingest 시 UPSERT 로 행 수가 늘지 않음.
  4. embed 실패 시 fail-soft — errors 수집, 예외 전파 안 함.

fixture: .pptx 는 python-pptx 로 tmp_path 에 런타임 생성(바이너리 커밋 금지).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from core.ingest.parser_base import ParserRegistry
from core.ingest.parsers.pptx import PptxParser
from core.ingest.pipeline import INGEST_SOURCE, DocumentIngestPipeline
from core.rag.knowledge_store import KnowledgeStore


# ─────────────────────────────────────────────
# Fake 모델 프로바이더 — embed 만 구현
# ─────────────────────────────────────────────
class FakeEmbedProvider:
    """
    embed(texts) 만 구현한 가짜 프로바이더.

    파이프라인은 ModelProvider.embed 만 사용하므로, 다른 메서드는 필요 없다.
    각 텍스트마다 고정 차원(dim)의 더미 벡터를 반환한다(내용 무관, 차원만 일정).
    """

    def __init__(self, dim: int = 8) -> None:
        self.dim = dim
        # 호출된 텍스트를 기록해 passage 접두사 등을 검증할 수 있게 한다.
        self.calls: list[list[str]] = []

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls.append(list(texts))
        # 텍스트별로 살짝 다른 값(인덱스 기반)을 줘 동일 벡터 충돌을 피한다.
        return [[float(i % 7) + 0.1 for _ in range(self.dim)] for i in range(len(texts))]


class BrokenEmbedProvider:
    """embed 호출 시 RuntimeError 를 던지는 가짜 프로바이더(fail-soft 검증용)."""

    async def embed(self, texts: list[str]) -> list[list[float]]:
        raise RuntimeError("embedding server down")


# ─────────────────────────────────────────────
# 헬퍼 — 두 소제목(슬라이드)을 가진 .pptx 생성
# ─────────────────────────────────────────────
def _make_two_topic_pptx(tmp_path: Path) -> Path:
    """
    '보안' / '네트워크' 두 슬라이드 + 각 본문을 가진 .pptx 를 만든다.

    충돌 방지 검증과 metadata 보존 검증을 한 파일로 처리한다.
    """
    prs = Presentation()

    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "보안"
    slide1.placeholders[1].text_frame.text = "접근 통제는 최소 권한 원칙을 따른다."

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "네트워크"
    slide2.placeholders[1].text_frame.text = "방화벽은 기본 차단 정책을 적용한다."

    path = tmp_path / "policy.pptx"
    prs.save(str(path))
    return path


def _build_registry() -> ParserRegistry:
    """PPTX 파서가 등록된 레지스트리를 만든다."""
    reg = ParserRegistry()
    reg.register(PptxParser())
    return reg


# ─────────────────────────────────────────────
# 1) dry_run=True — 적재 없이 요약
# ─────────────────────────────────────────────
@pytest.mark.asyncio
async def test_ingest_dry_run_summarizes_without_storing(tmp_path: Path) -> None:
    """dry_run=True 면 KnowledgeStore 에 아무것도 적재하지 않고 요약만 돌려준다."""
    path = _make_two_topic_pptx(tmp_path)
    store = KnowledgeStore(pg_pool=None)
    embedder = FakeEmbedProvider()
    pipeline = DocumentIngestPipeline(
        parser_registry=_build_registry(),
        model_provider=embedder,
        knowledge_store=store,
    )

    summary = await pipeline.ingest_file(path, dry_run=True)

    assert summary["dry_run"] is True
    assert summary["ingested"] == 0
    assert summary["doc_format"] == "pptx"
    assert summary["title"] == "보안"  # 첫 슬라이드 제목
    assert summary["chunk_count"] > 0
    assert summary["parent_count"] >= 2  # 두 그룹(보안/네트워크) → 부모 2개 이상
    # 적재가 일어나지 않았으므로 임베딩도 호출되지 않아야 한다.
    assert embedder.calls == []
    assert await store.count() == 0


# ─────────────────────────────────────────────
# 2) dry_run=False — 적재 + metadata/embedding 보존
# ─────────────────────────────────────────────
@pytest.mark.asyncio
async def test_ingest_stores_entries_with_metadata(tmp_path: Path) -> None:
    """
    dry_run=False 시 KnowledgeEntry 가 적재되고, source/metadata/embedding 이
    올바르게 보존되는지 검증한다.
    """
    path = _make_two_topic_pptx(tmp_path)
    store = KnowledgeStore(pg_pool=None)
    embedder = FakeEmbedProvider(dim=8)
    pipeline = DocumentIngestPipeline(
        parser_registry=_build_registry(),
        model_provider=embedder,
        knowledge_store=store,
    )

    summary = await pipeline.ingest_file(path, dry_run=False)

    assert summary["dry_run"] is False
    assert summary["ingested"] == summary["chunk_count"]
    assert summary["errors"] == []
    assert await store.count(INGEST_SOURCE) == summary["chunk_count"]

    # 적재된 엔트리 검사 — 인메모리 store._store 에 KnowledgeEntry 가 들어 있다.
    entries = list(store._store.values())
    assert entries, "엔트리가 적재돼야 한다"
    for e in entries:
        assert e.source == INGEST_SOURCE
        # embedding 은 tuple 이어야 한다(KnowledgeStore.add 계약).
        assert isinstance(e.embedding, tuple)
        assert len(e.embedding) == 8
        # metadata 구조 메타 보존.
        assert "element_type" in e.metadata
        assert "heading_path" in e.metadata
        assert "page" in e.metadata
        assert e.metadata["doc_format"] == "pptx"

    # 임베딩 요청은 "passage: " 접두사를 붙여 보낸다(e5-large 규약).
    assert embedder.calls, "임베딩이 호출돼야 한다"
    assert all(t.startswith("passage: ") for batch in embedder.calls for t in batch)


@pytest.mark.asyncio
async def test_ingest_preserves_heading_path_per_topic(tmp_path: Path) -> None:
    """
    적재된 엔트리의 heading_path 가 슬라이드별로 분리 보존되는지 — 충돌 0 의 결과.

    '보안' 본문 엔트리와 '네트워크' 본문 엔트리의 heading_path 가 섞이지 않아야 한다.
    """
    path = _make_two_topic_pptx(tmp_path)
    store = KnowledgeStore(pg_pool=None)
    pipeline = DocumentIngestPipeline(
        parser_registry=_build_registry(),
        model_provider=FakeEmbedProvider(),
        knowledge_store=store,
    )
    await pipeline.ingest_file(path, dry_run=False)

    # content 로 슬라이드를 식별해 heading_path 가 올바르게 매칭되는지 확인.
    for e in store._store.values():
        hp = tuple(e.metadata["heading_path"])
        if "최소 권한" in e.content:
            assert hp == ("보안",)
        if "방화벽" in e.content:
            assert hp == ("네트워크",)
        # 두 토픽 텍스트가 한 엔트리에 섞이면 안 된다(충돌 0).
        assert not ("최소 권한" in e.content and "방화벽" in e.content)


# ─────────────────────────────────────────────
# 3) 멱등성 — 2회 ingest 시 행 수 불변
# ─────────────────────────────────────────────
@pytest.mark.asyncio
async def test_ingest_is_idempotent_on_reingest(tmp_path: Path) -> None:
    """같은 파일을 두 번 적재해도 UPSERT 로 행 수가 늘지 않아야 한다."""
    path = _make_two_topic_pptx(tmp_path)
    store = KnowledgeStore(pg_pool=None)
    pipeline = DocumentIngestPipeline(
        parser_registry=_build_registry(),
        model_provider=FakeEmbedProvider(),
        knowledge_store=store,
    )

    first = await pipeline.ingest_file(path, dry_run=False)
    count_after_first = await store.count(INGEST_SOURCE)

    second = await pipeline.ingest_file(path, dry_run=False)
    count_after_second = await store.count(INGEST_SOURCE)

    # 두 번째 적재도 같은 청크 수를 처리하지만, id 가 동일해 행 수는 그대로다.
    assert first["chunk_count"] == second["chunk_count"]
    assert count_after_first == count_after_second


# ─────────────────────────────────────────────
# 4) embed 실패 — fail-soft (errors 수집, 예외 전파 없음)
# ─────────────────────────────────────────────
@pytest.mark.asyncio
async def test_ingest_embed_failure_is_failsoft(tmp_path: Path) -> None:
    """
    임베딩이 예외를 던져도 ingest_file 은 예외를 전파하지 않고 errors 에 모은다.

    적재(ingested)는 0 이지만, 파일 단위 처리는 중단되지 않아야 한다.
    """
    path = _make_two_topic_pptx(tmp_path)
    store = KnowledgeStore(pg_pool=None)
    pipeline = DocumentIngestPipeline(
        parser_registry=_build_registry(),
        model_provider=BrokenEmbedProvider(),
        knowledge_store=store,
    )

    # 예외가 전파되지 않아야 한다.
    summary = await pipeline.ingest_file(path, dry_run=False)

    assert summary["ingested"] == 0
    assert summary["errors"], "임베딩 실패가 errors 에 기록돼야 한다"
    assert any("임베딩 실패" in msg for msg in summary["errors"])
    # 적재된 행이 없어야 한다.
    assert await store.count(INGEST_SOURCE) == 0


# ─────────────────────────────────────────────
# 5) 미지원 확장자 — fail-closed
# ─────────────────────────────────────────────
@pytest.mark.asyncio
async def test_ingest_unsupported_extension_returns_error_summary(tmp_path: Path) -> None:
    """지원 파서가 없는 확장자는 빈 요약 + errors 를 반환한다(fail-closed)."""
    txt = tmp_path / "note.txt"
    txt.write_text("그냥 텍스트", encoding="utf-8")
    store = KnowledgeStore(pg_pool=None)
    pipeline = DocumentIngestPipeline(
        parser_registry=_build_registry(),
        model_provider=FakeEmbedProvider(),
        knowledge_store=store,
    )

    summary = await pipeline.ingest_file(txt, dry_run=False)

    assert summary["chunk_count"] == 0
    assert summary["ingested"] == 0
    assert summary["errors"]
    assert any("지원하는 파서가 없습니다" in m for m in summary["errors"])
