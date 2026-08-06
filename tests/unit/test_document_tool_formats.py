# DocumentProcess 도구의 포맷 커버리지(PPTX/HWPX/평문/바이너리 거부)를 검증하는 단위 테스트.
"""DocumentProcess 포맷 확장(W8 업로드 확장) 단위 테스트.

무엇을 지키나:
    - .pptx / .hwpx 가 core/ingest 파서로 라우팅돼 실제 본문 텍스트가 나온다.
      (예전에는 UTF-8 평문으로 읽혀 깨진 문자열이 모델에 들어갔다 — 이 회귀를 막는다.)
    - 확장자만 .hwp 이고 내용이 다른 파일은 조용히 쓰레기를 내지 않고 오류로 거절한다
      (파서 can_parse 의 매직바이트 검사 = fail-closed).
    - 알 수 없는 확장자라도 평문이면 그대로 읽고, 바이너리면 명확한 오류를 낸다.
    - 한글 CP949(euc-kr) 텍스트가 물음표로 뭉개지지 않는다.
    - 개행이 CRLF 인 파일도 LF 로 정규화돼 글자 수(=청크 경계)가 흔들리지 않는다.

테스트 전략(.claude/rules/testing.md 준수):
    - 파일시스템은 tmp_path fixture 만 사용한다(프로젝트 파일 수정 금지).
    - 픽스처는 python-pptx / python-hwpx 로 테스트 실행 시점에 만든다. 두 라이브러리는
      DocumentExport 도구도 쓰는 기존 의존성이며, 없으면 해당 테스트만 skip 한다.
    - .hwp 구포맷은 LibreOffice(soffice) 바이너리가 필요해 CI 에서 재현할 수 없으므로,
      "변환 없이도 확정적으로 판정되는" 매직바이트 거부 경로만 검증한다.
    - asyncio_mode="auto" 라 async def test_* 는 데코레이터 없이 동작한다.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from core.tools.base import ToolUseContext
from core.tools.implementations.document_tool import DocumentProcessTool


def _ctx() -> ToolUseContext:
    """기본 옵션(예산 미주입)으로 도구 컨텍스트를 만든다."""
    return ToolUseContext(cwd=".", options={})


async def _run(path: Path):
    """DocumentProcess를 한 번 호출해 ToolResult를 돌려주는 헬퍼."""
    return await DocumentProcessTool().call({"file_path": str(path)}, _ctx())


# ─────────────────────────────────────────────
# 픽스처 생성 — 실제 바이너리 문서를 만든다
# ─────────────────────────────────────────────


def _make_pptx(path: Path) -> None:
    """슬라이드 2장짜리 .pptx 를 만든다(제목 + 본문 + 텍스트박스)."""
    pptx = pytest.importorskip("pptx", reason="python-pptx 미설치")
    from pptx.util import Inches

    prs = pptx.Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "업로드 확장 검증"
    slide1.placeholders[1].text = "첫 번째 슬라이드 본문"

    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "두 번째 슬라이드"
    box = slide2.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    box.text_frame.text = "텍스트박스 안의 한글"

    prs.save(str(path))


def _make_hwpx(path: Path) -> None:
    """제목 + 단락 2개짜리 .hwpx 를 만든다."""
    builder = pytest.importorskip("hwpx.builder", reason="python-hwpx 미설치")

    doc = builder.Document(
        sections=[
            builder.Section(
                children=[
                    builder.Heading(level=1, text="한글 문서 제목"),
                    builder.Paragraph(text="한글 본문 첫 단락"),
                    builder.Paragraph(text="한글 본문 둘째 단락"),
                ]
            )
        ]
    )
    doc.save_to_path(str(path))


# ─────────────────────────────────────────────
# PPTX / HWPX — 인제스트 파서 라우팅
# ─────────────────────────────────────────────


async def test_document_process_pptx_extracts_slide_text(tmp_path: Path):
    """.pptx 는 슬라이드 제목·본문·텍스트박스 텍스트를 모두 뽑아낸다."""
    path = tmp_path / "deck.pptx"
    _make_pptx(path)

    result = await _run(path)

    assert not result.is_error, result.error_message
    assert "업로드 확장 검증" in result.data
    assert "첫 번째 슬라이드 본문" in result.data
    assert "텍스트박스 안의 한글" in result.data


async def test_document_process_hwpx_extracts_paragraph_text(tmp_path: Path):
    """.hwpx 는 제목과 본문 단락을 텍스트로 뽑아낸다(한글 세그먼트 핵심 포맷)."""
    path = tmp_path / "doc.hwpx"
    _make_hwpx(path)

    result = await _run(path)

    assert not result.is_error, result.error_message
    assert "한글 문서 제목" in result.data
    assert "한글 본문 첫 단락" in result.data
    assert "한글 본문 둘째 단락" in result.data


async def test_document_process_hwp_empty_result_gives_actionable_hint(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """.hwp 변환이 실패하면 내부 경고가 아니라 'HWPX/PDF로 저장하라'는 안내가 나온다.

    배경(2026-08-06 실측): LibreOffice 한글 필터는 HWP V2/V3(한글 97 이하)만
    인식하고 한글 2002 이후의 HWP 5.0은 지원하지 않는다. 그래서 실사용 .hwp는
    변환이 비게 되는데, 이때 "soffice 미설치" 같은 내부 경고만 보이면 사용자는
    같은 파일을 계속 다시 올린다. 조치 가능한 안내가 나와야 한다.
    """
    from core.ingest.types import DocumentTree
    from core.tools.implementations import document_tool

    class _EmptyParser:
        """빈 트리 + 경고를 돌려주는 파서(변환 실패 상황 재현)."""

        async def parse(self, path: Path) -> DocumentTree:
            return DocumentTree(
                title=path.stem,
                source_path=str(path),
                nodes=(),
                doc_format="hwp",
                warnings=("soffice 실행 파일을 찾지 못했습니다",),
            )

    class _Registry:
        def get_for_path(self, path: Path):
            return _EmptyParser()

    monkeypatch.setattr(document_tool, "_get_ingest_registry", lambda: _Registry())

    path = tmp_path / "old.hwp"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)

    result = await _run(path)

    assert not result.is_error
    assert "HWPX" in result.data and "PDF" in result.data
    # 진단은 남기되 주가 되어선 안 된다 — 안내가 먼저 온다.
    assert result.data.index("HWPX") < result.data.index("soffice")


async def test_document_process_fake_hwp_rejected_not_garbled(tmp_path: Path):
    """확장자만 .hwp 인 파일은 쓰레기 텍스트가 아니라 오류로 거절된다(fail-closed)."""
    path = tmp_path / "fake.hwp"
    # OLE2 시그니처가 아닌 바이트 → HwpViaLibreOfficeParser.can_parse 가 False.
    path.write_bytes(b"\xff\xfe\x00\x01not really hwp\x00")

    result = await _run(path)

    assert result.is_error
    assert ".hwp" in result.error_message


# ─────────────────────────────────────────────
# 평문 경로 — 인코딩/개행/바이너리 판별
# ─────────────────────────────────────────────


async def test_document_process_unknown_binary_returns_clear_error(tmp_path: Path):
    """모르는 확장자라도 바이너리면 깨진 텍스트 대신 지원 형식을 안내한다."""
    path = tmp_path / "blob.bin"
    path.write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00")

    result = await _run(path)

    assert result.is_error
    assert "지원하지 않는 파일 형식" in result.error_message
    assert "HWPX" in result.error_message  # 지원 목록을 함께 알려준다


async def test_document_process_cp949_text_decoded_not_mojibake(tmp_path: Path):
    """UTF-8이 아닌 CP949 한글 텍스트도 물음표로 뭉개지지 않고 그대로 읽힌다."""
    path = tmp_path / "legacy.txt"
    path.write_bytes("한글 레거시 인코딩 문서".encode("cp949"))

    result = await _run(path)

    assert not result.is_error, result.error_message
    assert "한글 레거시 인코딩 문서" in result.data
    assert "�" not in result.data  # 대체 문자(U+FFFD)가 섞이면 실패


async def test_document_process_utf16_bom_decoded_not_rejected(tmp_path: Path):
    """UTF-16 텍스트(메모장 '유니코드' 저장)는 NUL이 많아도 바이너리로 오거부되지 않는다."""
    path = tmp_path / "unicode.txt"
    path.write_bytes("한글 UTF-16 문서".encode("utf-16"))  # BOM 포함

    result = await _run(path)

    assert not result.is_error, result.error_message
    assert "한글 UTF-16 문서" in result.data


async def test_document_process_utf8_bom_stripped(tmp_path: Path):
    """UTF-8 BOM은 본문 앞에 보이지 않는 문자로 남지 않는다."""
    path = tmp_path / "bom.txt"
    path.write_bytes(b"\xef\xbb\xbf" + "본문 시작".encode())

    result = await _run(path)

    assert not result.is_error, result.error_message
    assert "﻿" not in result.data
    assert result.data.count("본문 시작") == 1


async def test_document_process_crlf_normalized_to_lf(tmp_path: Path):
    """CRLF 파일도 LF로 정규화돼 글자 수(=청크 경계 계산)가 부풀지 않는다."""
    path = tmp_path / "crlf.txt"
    path.write_bytes(b"first line\r\nsecond line\r\n")

    result = await _run(path)

    assert not result.is_error, result.error_message
    assert "\r" not in result.data
    # "first line\nsecond line\n" = 23자 — \r 이 남아 있었다면 25자가 된다.
    assert result.metadata["total_chars"] == 23
