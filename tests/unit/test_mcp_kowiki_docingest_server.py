"""
mcp_servers/{kowiki,docingest,diag}_server.py 도구 단위 테스트.

운영 DB 오염 방지 (절대 준수):
  - 실 임베딩 서버(:8002)/실 PG/실 GPU 미접근.
  - 임베딩은 _FakeEmbedProvider(고정 벡터), 저장소는 인메모리 KnowledgeStore(pg=None).
  - diag 는 paramiko/socket/urllib 을 patch 하여 실 네트워크 차단.

검증 대상:
  kowiki.KowikiSearchTool
    - "query: " 접두사로 임베딩하는지, 결과 형식(content/metadata/score) 정합.
    - query 누락/빈 문자열 → ValueError, top_k 검증/상한.
  docingest.{DocParseTool, DocIngestTool, DocSearchTool}
    - python-pptx 로 만든 .pptx fixture 로 parse(dry_run)/ingest(쓰기) 동작.
    - ingest 후 인메모리 store 에서 search 로 회수되는지(전 구간).
  diag.ReachabilityTool / RagLatencyTool
    - reachability: 하위 점검 함수를 patch → 구조화 JSON({web,db,gpu}) 반환.
    - rag_latency: 임베딩 호출을 patch, pg_pool=None 이면 search skipped 반환.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest
from pptx import Presentation

from core.ingest.parser_base import ParserRegistry
from core.ingest.parsers.pptx import PptxParser
from core.ingest.pipeline import DocumentIngestPipeline
from core.rag.knowledge_store import KnowledgeEntry, KnowledgeStore
from mcp_servers.docingest_server import (
    DocIngestTool,
    DocParseTool,
    DocSearchTool,
)
from mcp_servers.kowiki_server import _QUERY_PREFIX, KowikiSearchTool


# ─────────────────────────────────────────────
# 가짜 임베딩 프로바이더 — 고정/결정론적 벡터
# ─────────────────────────────────────────────
class _FakeEmbedProvider:
    """
    embed(texts) -> list[list[float]] 만 제공하는 가짜 프로바이더.

    결정론을 위해 텍스트 길이 기반의 단순 벡터를 만든다. 차원은 작아도
    인메모리 코사인 검색에는 충분하다(실 1024차원과 무관 — 폴백 경로 검증).
    """

    def __init__(self, dim: int = 8) -> None:
        self._dim = dim
        # 검증용 — 마지막으로 임베딩한 텍스트 목록을 기록한다.
        self.last_texts: list[str] | None = None

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.last_texts = list(texts)
        out: list[list[float]] = []
        for t in texts:
            # 첫 성분에 길이를 넣고 나머지는 1.0 — 모든 청크가 검색에 걸리도록.
            vec = [float(len(t) % 7 + 1)] + [1.0] * (self._dim - 1)
            out.append(vec)
        return out


# ─────────────────────────────────────────────
# kowiki.KowikiSearchTool
# ─────────────────────────────────────────────
class TestKowikiSearchTool:
    """tb_knowledge 벡터 검색 도구(kowiki) — 인메모리 폴백."""

    async def test_search_uses_query_prefix(self):
        """질의를 'query: ' 접두사로 임베딩해야 한다(e5 규약)."""
        provider = _FakeEmbedProvider()
        store = KnowledgeStore(pg_pool=None)  # 인메모리 폴백
        tool = KowikiSearchTool(provider, store)

        await tool.call({"query": "성리학"})

        assert provider.last_texts == [f"{_QUERY_PREFIX}성리학"]

    async def test_search_returns_results_with_expected_shape(self):
        """검색 결과가 content/metadata/score(similarity) 형식을 가져야 한다."""
        provider = _FakeEmbedProvider()
        store = KnowledgeStore(pg_pool=None)
        # 인메모리에 청크 1건 적재(임베딩 포함).
        await store.add(
            KnowledgeEntry(
                source="kowiki",
                title="성리학",
                content="이기이원론 설명",
                section="개요",
                embedding=tuple([1.0] * 8),
            )
        )
        tool = KowikiSearchTool(provider, store)

        result = await tool.call({"query": "이기이원론", "top_k": 3})

        assert result["query"] == "이기이원론"
        assert result["count"] >= 1
        first = result["results"][0]
        # KnowledgeStore.search_by_vector 결과 키 정합.
        assert "content" in first
        assert "similarity" in first  # score
        assert "metadata" in first
        assert first["content"] == "이기이원론 설명"

    async def test_search_empty_query_raises_value_error(self):
        """빈/공백 query 는 ValueError 여야 한다."""
        tool = KowikiSearchTool(_FakeEmbedProvider(), KnowledgeStore(pg_pool=None))
        with pytest.raises(ValueError):
            await tool.call({"query": "   "})

    async def test_search_missing_query_raises_value_error(self):
        """query 누락 → ValueError."""
        tool = KowikiSearchTool(_FakeEmbedProvider(), KnowledgeStore(pg_pool=None))
        with pytest.raises(ValueError):
            await tool.call({})

    async def test_search_negative_top_k_raises_value_error(self):
        """음수 top_k 는 ValueError 여야 한다(1 미만 거부)."""
        tool = KowikiSearchTool(_FakeEmbedProvider(), KnowledgeStore(pg_pool=None))
        with pytest.raises(ValueError):
            await tool.call({"query": "x", "top_k": -3})

    async def test_search_zero_top_k_raises_value_error(self):
        """top_k=0 은 거부되어야 한다(v7.3 견고화 — falsy 치환 폐기).

        과거 `arguments.get('top_k') or _DEFAULT_TOP_K` 관용구는 0 을 falsy 로
        보아 기본값(5)으로 조용히 치환했다. 이제는 `get(키, 기본값)` + 정수/범위
        검증으로 바뀌어, 0 은 '잘못된 입력'으로 ValueError(프레임워크 -32602)가 된다.
        """
        tool = KowikiSearchTool(_FakeEmbedProvider(), KnowledgeStore(pg_pool=None))
        with pytest.raises(ValueError):
            await tool.call({"query": "x", "top_k": 0})

    async def test_search_bool_top_k_raises_value_error(self):
        """top_k=True(bool)는 거부되어야 한다.

        파이썬에서 bool 은 int 의 하위형이라 isinstance(True, int) 가 True 다.
        견고화 코드는 `isinstance(top_k, bool)` 가드로 bool 을 명시적으로 배제한다
        (True 가 1 로 통과해 의미가 모호해지는 것을 막는다).
        """
        tool = KowikiSearchTool(_FakeEmbedProvider(), KnowledgeStore(pg_pool=None))
        with pytest.raises(ValueError):
            await tool.call({"query": "x", "top_k": True})

    async def test_search_default_top_k_when_unspecified(self):
        """top_k 미지정 시 기본값(_DEFAULT_TOP_K=5)으로 정상 동작해야 한다."""
        tool = KowikiSearchTool(_FakeEmbedProvider(), KnowledgeStore(pg_pool=None))
        # 예외 없이 정상 반환(미지정 → 기본 5, 거부 아님).
        result = await tool.call({"query": "x"})
        assert result["query"] == "x"
        assert result["count"] == 0  # 인메모리 빈 store

    async def test_search_top_k_clamped_to_max(self):
        """top_k 가 상한(_MAX_TOP_K=50)을 넘으면 클램프되어 검색이 수행돼야 한다.

        store 가 top_k 를 그대로 검색에 넘기므로, 거대한 top_k(100)도 예외 없이
        50 으로 잘려 정상 동작한다. 검증은 '예외 없이 반환' 으로 충분하다
        (인메모리 store 결과 수는 적재량에 종속이라 상한 자체를 직접 보긴 어렵다)."""
        tool = KowikiSearchTool(_FakeEmbedProvider(), KnowledgeStore(pg_pool=None))
        result = await tool.call({"query": "x", "top_k": 100})
        assert result["query"] == "x"

    async def test_search_string_top_k_raises_value_error(self):
        """top_k 가 정수가 아닌 문자열이면 ValueError(타입 가드)."""
        tool = KowikiSearchTool(_FakeEmbedProvider(), KnowledgeStore(pg_pool=None))
        with pytest.raises(ValueError):
            await tool.call({"query": "x", "top_k": "5"})

    async def test_search_embed_failure_normalized_to_runtime_error(self):
        """임베딩 실패는 RuntimeError 로 정규화되어야 한다(프레임워크 -32603)."""

        class _BoomProvider:
            async def embed(self, texts: list[str]) -> list[list[float]]:
                raise OSError("임베딩 서버 다운")

        tool = KowikiSearchTool(_BoomProvider(), KnowledgeStore(pg_pool=None))
        with pytest.raises(RuntimeError):
            await tool.call({"query": "x"})

    async def test_search_empty_embedding_raises_runtime_error(self):
        """임베딩 서버가 빈 결과를 주면 RuntimeError."""

        class _EmptyProvider:
            async def embed(self, texts: list[str]) -> list[list[float]]:
                return [[]]

        tool = KowikiSearchTool(_EmptyProvider(), KnowledgeStore(pg_pool=None))
        with pytest.raises(RuntimeError):
            await tool.call({"query": "x"})


# ─────────────────────────────────────────────
# docingest — parse/ingest/search 전 구간 (인메모리)
# ─────────────────────────────────────────────
def _make_pptx(tmp_path: Path) -> Path:
    """python-pptx 로 텍스트가 있는 .pptx fixture 를 런타임 생성한다(바이너리 미커밋)."""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    box = slide.shapes.add_textbox(0, 0, 4_000_000, 1_000_000)
    box.text_frame.text = "성리학의 이기이원론은 이와 기를 구분한다."
    path = tmp_path / "doc.pptx"
    prs.save(str(path))
    return path


def _make_pipeline() -> tuple[DocumentIngestPipeline, _FakeEmbedProvider, KnowledgeStore]:
    """fake embed + 인메모리 store 로 파이프라인을 조립한다(실 DB 미접근)."""
    registry = ParserRegistry()
    registry.register(PptxParser(), priority=0)
    provider = _FakeEmbedProvider()
    store = KnowledgeStore(pg_pool=None)
    pipeline = DocumentIngestPipeline(
        parser_registry=registry,
        model_provider=provider,  # type: ignore[arg-type]
        knowledge_store=store,
    )
    return pipeline, provider, store


class TestDocParseTool:
    """parse 도구 — read-only(dry_run): 적재 없이 청킹 요약."""

    async def test_parse_dry_run_returns_summary_without_ingesting(self, tmp_path: Path):
        """parse 는 dry_run=True 라 ingested=0 이고 store 에 아무것도 적재되지 않아야 한다."""
        pipeline, _provider, store = _make_pipeline()
        tool = DocParseTool(pipeline)
        pptx = _make_pptx(tmp_path)

        result = await tool.call({"path": str(pptx)})

        assert result["dry_run"] is True
        assert result["ingested"] == 0
        assert result["chunk_count"] >= 1
        # 인메모리 store 가 비어 있어야 한다(부작용 없음).
        assert await store.count() == 0

    async def test_parse_missing_path_raises_value_error(self):
        """path 누락 → ValueError."""
        pipeline, _p, _s = _make_pipeline()
        tool = DocParseTool(pipeline)
        with pytest.raises(ValueError):
            await tool.call({})

    async def test_parse_nonexistent_file_returns_fail_soft_summary(self):
        """존재하지 않는 .pptx 는 예외가 아니라 fail-soft 요약(errors 채움)을 반환한다.

        PptxParser.can_parse 가 매직바이트(ZIP) 확인에 실패해 None 파서가 되고,
        파이프라인은 errors 에 '지원하는 파서가 없습니다'를 담아 반환한다(예외 없음).
        DocParseTool 은 이 요약을 그대로 돌려준다.
        """
        pipeline, _p, _s = _make_pipeline()
        tool = DocParseTool(pipeline)
        result = await tool.call({"path": "/no/such/file.pptx"})
        assert result["chunk_count"] == 0
        assert result["errors"]  # 사유가 채워져야 한다


class TestDocIngestAndSearch:
    """ingest(쓰기) 후 search(read-only)로 회수되는 전 구간을 검증한다."""

    async def test_ingest_then_search_round_trip(self, tmp_path: Path):
        """ingest 로 적재한 문서를 search 로 다시 찾을 수 있어야 한다(인메모리)."""
        pipeline, provider, store = _make_pipeline()
        pptx = _make_pptx(tmp_path)

        # 1) 적재(쓰기).
        ingest_tool = DocIngestTool(pipeline)
        ingest_result = await ingest_tool.call({"path": str(pptx)})
        assert ingest_result["dry_run"] is False
        assert ingest_result["ingested"] >= 1
        # 임베딩은 "passage: " 접두사로 적재되어야 한다(인제스트 규약).
        assert provider.last_texts is not None
        assert all(t.startswith("passage: ") for t in provider.last_texts)
        # store 에 실제로 행이 들어갔는지.
        assert await store.count() >= 1

        # 2) 검색(read-only) — source='docingest' 필터로 회수.
        search_tool = DocSearchTool(provider, store)
        search_result = await search_tool.call({"query": "이기이원론"})
        assert search_result["count"] >= 1
        # 적재된 본문 일부가 결과에 나타나야 한다.
        contents = " ".join(r["content"] for r in search_result["results"])
        assert "이기이원론" in contents

    async def test_ingest_missing_path_raises_value_error(self):
        """ingest path 누락 → ValueError."""
        pipeline, _p, _s = _make_pipeline()
        tool = DocIngestTool(pipeline)
        with pytest.raises(ValueError):
            await tool.call({})

    async def test_search_filters_by_docingest_source(self, tmp_path: Path):
        """DocSearchTool 은 source='docingest' 로 필터해 다른 소스를 누설하지 않아야 한다."""
        pipeline, provider, store = _make_pipeline()
        # 다른 소스(kowiki)로 청크를 하나 심어 둔다 — 검색에 걸리면 안 됨.
        await store.add(
            KnowledgeEntry(
                source="kowiki",
                title="다른 소스",
                content="이것은 kowiki 소스 문서이다 이기이원론",
                embedding=tuple([1.0] * 8),
            )
        )
        # docingest 소스 문서를 적재.
        await DocIngestTool(pipeline).call({"path": str(_make_pptx(tmp_path))})

        result = await DocSearchTool(provider, store).call({"query": "이기이원론"})
        # 결과의 모든 항목이 docingest 소스여야 한다(kowiki 누설 없음).
        assert result["count"] >= 1
        assert all(r["source"] == "docingest" for r in result["results"])

    async def test_search_missing_query_raises_value_error(self):
        """DocSearchTool query 누락 → ValueError."""
        _pipeline, provider, store = _make_pipeline()
        tool = DocSearchTool(provider, store)
        with pytest.raises(ValueError):
            await tool.call({})

    @pytest.mark.parametrize("bad_top_k", [0, -3, True, "5"])
    async def test_docsearch_invalid_top_k_raises_value_error(self, bad_top_k):
        """DocSearchTool 도 top_k 견고화를 공유한다: 0/음수/bool/문자열 → ValueError.

        kowiki 와 동일한 검증 로직(get(키, 기본값) + isinstance(int) + bool 가드 +
        <1 거부)을 docingest 검색에도 적용한다."""
        _pipeline, provider, store = _make_pipeline()
        tool = DocSearchTool(provider, store)
        with pytest.raises(ValueError):
            await tool.call({"query": "x", "top_k": bad_top_k})

    async def test_docsearch_top_k_clamped_to_max(self):
        """DocSearchTool: top_k=100 은 _MAX_TOP_K(=50)으로 클램프되어 store 에 전달돼야 한다.

        store 를 spy 로 감싸 search_by_vector 에 실제로 넘어간 top_k 가 50 인지
        직접 확인한다(상한 클램프의 강한 검증)."""
        from unittest.mock import AsyncMock

        _pipeline, provider, store = _make_pipeline()
        spy = AsyncMock(return_value=[])
        store.search_by_vector = spy  # type: ignore[assignment]
        tool = DocSearchTool(provider, store)

        await tool.call({"query": "x", "top_k": 100})

        # top_k 는 키워드 인자로 전달된다(코드: search_by_vector(emb, top_k=..., source=...)).
        assert spy.await_args.kwargs["top_k"] == 50


# ─────────────────────────────────────────────
# diag — reachability / rag_latency (네트워크 patch)
# ─────────────────────────────────────────────
class TestDiagReachabilityTool:
    """reachability 가 구조화 JSON 을 반환하는지 — 하위 점검 함수는 patch."""

    async def test_reachability_returns_structured_json(self):
        """web/db/gpu 점검 함수를 patch 해 {web,db,gpu} 구조를 반환하는지 검증."""
        from mcp_servers.diag_server import ReachabilityTool

        fake_web = {"reachable": True, "status": 200}
        fake_db = {"host": "192.168.10.39", "ports": {"PostgreSQL": {"reachable": True}}}
        fake_gpu = {"reachable": True, "host": "192.168.21.112"}

        # 블로킹 점검 함수(실 네트워크)를 전부 가짜로 교체한다.
        with (
            patch("mcp_servers.diag_server._check_web_blocking", return_value=fake_web),
            patch("mcp_servers.diag_server._check_db_blocking", return_value=fake_db),
            patch("mcp_servers.diag_server._check_gpu_blocking", return_value=fake_gpu),
        ):
            result = await ReachabilityTool().call({})

        assert result == {"web": fake_web, "db": fake_db, "gpu": fake_gpu}

    async def test_reachability_gpu_unreachable_shape(self):
        """GPU SSH 실패(가짜)도 구조화 결과로 환원되어야 한다(예외 전파 없음)."""
        from mcp_servers.diag_server import ReachabilityTool

        with (
            patch("mcp_servers.diag_server._check_web_blocking", return_value={"reachable": False}),
            patch("mcp_servers.diag_server._check_db_blocking", return_value={"ports": {}}),
            patch(
                "mcp_servers.diag_server._check_gpu_blocking",
                return_value={"reachable": False, "error": "SSH 실패"},
            ),
        ):
            result = await ReachabilityTool().call({})

        assert result["gpu"]["reachable"] is False
        assert "error" in result["gpu"]


class TestDiagRagLatencyTool:
    """rag_latency — 임베딩 호출 patch, pg_pool=None 이면 search skipped."""

    async def test_rag_latency_without_pg_skips_search(self):
        """pg_pool=None 이면 임베딩 지연만 측정하고 search 는 skipped 여야 한다."""
        from mcp_servers.diag_server import RagLatencyTool

        # _embed_query_blocking 을 (벡터, 지연ms) 가짜로 교체(실 임베딩 서버 미접근).
        fake_vec = [0.1] * 1024
        with patch(
            "mcp_servers.diag_server._embed_query_blocking",
            return_value=(fake_vec, 12.3),
        ):
            tool = RagLatencyTool(pg_pool=None)
            result = await tool.call({"query": "이기이원론"})

        assert result["query"] == "이기이원론"
        assert result["embed_latency_ms"] == 12.3
        assert result["embed_dim"] == 1024
        assert "skipped" in result["search"]

    async def test_rag_latency_embed_failure_normalized_to_runtime_error(self):
        """임베딩 호출 실패는 RuntimeError 로 정규화되어야 한다."""
        from mcp_servers.diag_server import RagLatencyTool

        with patch(
            "mcp_servers.diag_server._embed_query_blocking",
            side_effect=OSError("연결 거부"),
        ):
            tool = RagLatencyTool(pg_pool=None)
            with pytest.raises(RuntimeError):
                await tool.call({"query": "x"})

    @pytest.mark.parametrize("bad_top_k", [0, -3, True, "5"])
    async def test_rag_latency_invalid_top_k_raises_value_error(self, bad_top_k):
        """rag_latency 도 top_k 견고화를 공유한다: 0/음수/bool/문자열 → ValueError.

        v7.3 견고화 전에는 top_k=0 이 falsy 라 기본값(5)으로 치환되어 거부되지
        않았다. 이제는 0 도 명시적으로 거부된다(get(키, 기본값) + 정수/범위 검증).
        검증이 임베딩 호출보다 먼저 일어나므로, _embed_query_blocking 을 patch 하지
        않아도 네트워크 접근 없이 ValueError 가 난다."""
        from mcp_servers.diag_server import RagLatencyTool

        tool = RagLatencyTool(pg_pool=None)
        with pytest.raises(ValueError):
            await tool.call({"top_k": bad_top_k})

    async def test_rag_latency_default_top_k_when_unspecified(self):
        """top_k 미지정 시 기본값(5)으로 정상 동작해야 한다(거부 아님)."""
        from mcp_servers.diag_server import RagLatencyTool

        with patch(
            "mcp_servers.diag_server._embed_query_blocking",
            return_value=([0.1] * 1024, 5.0),
        ):
            tool = RagLatencyTool(pg_pool=None)
            result = await tool.call({"query": "x"})  # top_k 미지정
        # 예외 없이 임베딩 지연을 측정해야 한다.
        assert result["embed_dim"] == 1024
        assert "skipped" in result["search"]
