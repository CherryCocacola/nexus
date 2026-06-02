"""
ParserRegistry / DocumentIngestPipeline 의 .pdf/.hwpx 라우팅 단위 테스트 (v7.3).

핵심 검증 의도:
  1. ParserRegistry 가 확장자별로 올바른 파서(PdfPlumberParser/HwpxParser/PptxParser)를
     반환하고, 미지원 확장자(.txt 등)에는 None 을 돌려준다(fail-closed).
  2. can_parse() fail-closed 와의 결합 — 확장자는 맞아도 매직바이트가 다르면 None.
  3. DocumentIngestPipeline.ingest_file(dry_run=True) 가 PDF/HWPX 를 적재 없이
     트리/청크 요약으로 돌려준다(외부 임베딩/DB 미접근 — fake provider + 인메모리 store).

외부 서비스 미접근:
  임베딩은 _FakeEmbedProvider(고정 벡터), 저장소는 KnowledgeStore(pg_pool=None)
  인메모리 폴백. 실 임베딩 서버/PG/GPU 에 접근하지 않는다.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib
import pytest
from hwpx.document import HwpxDocument

from core.ingest.parser_base import ParserRegistry
from core.ingest.parsers.hwpx import HwpxParser
from core.ingest.parsers.pdf_plumber import PdfPlumberParser
from core.ingest.parsers.pptx import PptxParser
from core.ingest.pipeline import DocumentIngestPipeline
from core.rag.knowledge_store import KnowledgeStore

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.backends.backend_pdf import PdfPages  # noqa: E402


# ─────────────────────────────────────────────
# 공용 fixture 헬퍼
# ─────────────────────────────────────────────
class _FakeEmbedProvider:
    """embed(texts)->list[list[float]] 만 제공하는 가짜 프로바이더(고정 차원 벡터)."""

    def __init__(self, dim: int = 8) -> None:
        self._dim = dim

    async def embed(self, texts: list[str]) -> list[list[float]]:
        # 모든 텍스트에 동일 차원의 1.0 벡터를 돌려준다(결정론).
        return [[1.0] * self._dim for _ in texts]


def _registry_with_all_parsers() -> ParserRegistry:
    """PPTX/PDF/HWPX 파서를 모두 등록한 레지스트리(docingest_server 와 동일 구성)."""
    reg = ParserRegistry()
    reg.register(PptxParser(), priority=0)
    reg.register(PdfPlumberParser(), priority=0)
    reg.register(HwpxParser(), priority=0)
    return reg


def _make_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "doc.pdf"
    with PdfPages(str(path)) as pdf:
        fig = plt.figure(figsize=(8.27, 11.69))
        fig.text(0.1, 0.85, "Body text for routing test.", fontsize=12)
        pdf.savefig(fig)
        plt.close(fig)
    return path


def _make_hwpx(tmp_path: Path) -> Path:
    doc = HwpxDocument.new()
    doc.add_paragraph("라우팅 테스트 본문 문단입니다.")
    path = tmp_path / "doc.hwpx"
    doc.save_to_path(str(path))
    return path


# ─────────────────────────────────────────────
# ParserRegistry 라우팅
# ─────────────────────────────────────────────
def test_registry_routes_pdf_to_pdf_parser(tmp_path: Path) -> None:
    """.pdf 경로는 PdfPlumberParser 로 라우팅돼야 한다."""
    reg = _registry_with_all_parsers()
    pdf = _make_pdf(tmp_path)
    parser = reg.get_for_path(pdf)
    assert isinstance(parser, PdfPlumberParser)


def test_registry_routes_hwpx_to_hwpx_parser(tmp_path: Path) -> None:
    """.hwpx 경로는 HwpxParser 로 라우팅돼야 한다."""
    reg = _registry_with_all_parsers()
    hwpx = _make_hwpx(tmp_path)
    parser = reg.get_for_path(hwpx)
    assert isinstance(parser, HwpxParser)


def test_registry_routes_pptx_to_pptx_parser(tmp_path: Path) -> None:
    """.pptx 경로는 PptxParser 로 라우팅돼야 한다(기존 파서 회귀)."""
    from pptx import Presentation

    reg = _registry_with_all_parsers()
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    pptx = tmp_path / "doc.pptx"
    prs.save(str(pptx))

    parser = reg.get_for_path(pptx)
    assert isinstance(parser, PptxParser)


def test_registry_unsupported_extension_returns_none(tmp_path: Path) -> None:
    """미지원 확장자(.txt)는 None 을 반환해야 한다(fail-closed — 임의 처리 금지)."""
    reg = _registry_with_all_parsers()
    txt = tmp_path / "doc.txt"
    txt.write_text("plain text", encoding="utf-8")
    assert reg.get_for_path(txt) is None


def test_registry_pdf_extension_wrong_magic_returns_none(tmp_path: Path) -> None:
    """확장자는 .pdf 지만 매직바이트가 아니면 can_parse 실패 → None.

    레지스트리는 후보 파서의 can_parse() 까지 통과해야 반환하므로, 가짜 PDF 는
    걸러진다(확장자만으로 처리하지 않는 fail-closed)."""
    reg = _registry_with_all_parsers()
    fake = tmp_path / "fake.pdf"
    fake.write_bytes(b"NOT-A-PDF")
    assert reg.get_for_path(fake) is None


def test_registry_supported_extensions_listing() -> None:
    """등록된 확장자 목록에 .hwpx/.pdf/.pptx 가 모두 포함되고 정렬돼야 한다."""
    reg = _registry_with_all_parsers()
    exts = reg.supported_extensions()
    assert set(exts) == {".hwpx", ".pdf", ".pptx"}
    assert list(exts) == sorted(exts)  # 정렬 보장


# ─────────────────────────────────────────────
# DocumentIngestPipeline dry_run 라우팅
# ─────────────────────────────────────────────
@pytest.mark.asyncio
async def test_pipeline_dry_run_pdf_returns_tree(tmp_path: Path) -> None:
    """파이프라인 dry_run 이 PDF 를 적재 없이 트리/청크 요약으로 반환해야 한다."""
    store = KnowledgeStore(pg_pool=None)
    pipeline = DocumentIngestPipeline(
        parser_registry=_registry_with_all_parsers(),
        model_provider=_FakeEmbedProvider(),  # type: ignore[arg-type]
        knowledge_store=store,
    )
    result = await pipeline.ingest_file(str(_make_pdf(tmp_path)), dry_run=True)

    assert result["dry_run"] is True
    assert result["doc_format"] == "pdf"
    assert result["ingested"] == 0
    assert result["chunk_count"] >= 1
    # dry_run 이므로 store 에 아무것도 적재되지 않아야 한다.
    assert await store.count() == 0


@pytest.mark.asyncio
async def test_pipeline_dry_run_hwpx_returns_tree(tmp_path: Path) -> None:
    """파이프라인 dry_run 이 HWPX 를 적재 없이 트리/청크 요약으로 반환해야 한다."""
    store = KnowledgeStore(pg_pool=None)
    pipeline = DocumentIngestPipeline(
        parser_registry=_registry_with_all_parsers(),
        model_provider=_FakeEmbedProvider(),  # type: ignore[arg-type]
        knowledge_store=store,
    )
    result = await pipeline.ingest_file(str(_make_hwpx(tmp_path)), dry_run=True)

    assert result["dry_run"] is True
    assert result["doc_format"] == "hwpx"
    assert result["ingested"] == 0
    assert result["chunk_count"] >= 1
    assert await store.count() == 0


@pytest.mark.asyncio
async def test_pipeline_unsupported_extension_returns_error_summary(tmp_path: Path) -> None:
    """미지원 확장자는 파서가 None 이라 errors 가 채워진 빈 요약을 반환한다(예외 없음)."""
    store = KnowledgeStore(pg_pool=None)
    pipeline = DocumentIngestPipeline(
        parser_registry=_registry_with_all_parsers(),
        model_provider=_FakeEmbedProvider(),  # type: ignore[arg-type]
        knowledge_store=store,
    )
    txt = tmp_path / "doc.txt"
    txt.write_text("plain", encoding="utf-8")

    result = await pipeline.ingest_file(str(txt), dry_run=True)
    assert result["chunk_count"] == 0
    assert result["errors"]
    assert "지원하는 파서가 없습니다" in result["errors"][0]
