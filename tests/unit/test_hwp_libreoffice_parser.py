"""
core/ingest/parsers/hwp_libreoffice.py — HwpViaLibreOfficeParser 단위 테스트.

대상: v7.3 로드맵 단계 9 — 구포맷 .hwp(한글 v5, OLE 복합문서).

검증 의도(여섯 갈래):
  1. can_parse — fail-closed: 확장자(.hwp) + OLE2 매직바이트까지 봐야 True.
     .hwpx/ZIP/잘못된 매직/없는 파일/확장자 불일치는 False.
  2. 변환 후 파싱(_parse_docx_structured) — 실 .docx fixture(python-docx 로 직접
     생성)를 파싱해 Heading1→SECTION / HeadingN→SUBHEADING / 일반→PARAGRAPH /
     표→TABLE, heading_path 누적, 문서 순서(order) 보존을 검증한다. mock 이 아니라
     실물 docx 를 쓴다(변환물 파싱 경로를 진짜로 태운다).
  3. _heading_level — 영문 "Heading N" / 한글 "제목 N" / Normal·Title(0) / 깊이 클램프.
  4. _table_node — " | " 직렬화, 빈 표 None, 셀 접근 실패 흡수.
  5. fail-soft — _convert_to_docx 의 subprocess.run 을 mock 해 FileNotFoundError /
     CalledProcessError / TimeoutExpired 주입 → warning + 빈 트리(예외 비전파).
     변환물 미발견 → warning.
  6. config/환경변수 — HwpConfig 기본값, NEXUS_SOFFICE_CMD 폴백, timeout 하한 보정.

왜 전 구간(실 soffice 변환→파싱) e2e 가 없는가:
  워크스페이스 전체에 실 .hwp 샘플이 없고, LibreOffice 는 HWP "export" 필터가
  없어 docx→hwp 생성도 불가하다. 따라서 두 절반을 분리 검증한다:
    · 변환 호출 단계: subprocess.run 을 mock 하여 fail-soft 만 확인.
    · 변환 후 파싱 단계: 실 .docx fixture 로 구조 보존 파싱을 확인(이 절반이
      실제로 LibreOffice 가 만들어줄 .docx 와 동형이므로, 파싱 정확성은 실물로 검증됨).

config 격리(왜 _cfg 를 직접 주입하거나 load_and_validate_config 를 patch 하는가):
  _ensure_config 는 첫 parse() 때 HwpConfig 를 읽어 self._cfg 에 캐시한다.
  변환/fail-soft 테스트는 실 config 의존을 피하려고 parser._cfg 를 미리 채워
  _ensure_config 를 단락(short-circuit)시킨다(FIRST 의 Independent).
"""

from __future__ import annotations

import subprocess
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document

from core.config import HwpConfig
from core.ingest.parsers.hwp_libreoffice import (
    _DEFAULT_SOFFICE_CMD,
    _DEFAULT_TIMEOUT_SEC,
    _HWP_OLE_MAGIC,
    HwpViaLibreOfficeParser,
)
from core.ingest.types import ElementType


# ─────────────────────────────────────────────
# fixture 생성 헬퍼 — 실 .hwp 매직 파일 / 실 .docx 문서
# ─────────────────────────────────────────────
def _write_hwp_magic(path: Path, *, valid: bool = True) -> Path:
    """can_parse 검증용 .hwp(OLE2 매직) 파일을 만든다.

    valid=True 면 진짜 OLE2 시그니처(D0CF11E0...)로 시작하는 파일을, False 면
    엉뚱한 바이트로 시작하는 파일을 쓴다. 본문은 의미 없으므로 매직 뒤 패딩만 둔다."""
    head = _HWP_OLE_MAGIC if valid else b"\x00\x01\x02\x03\x04\x05\x06\x07"
    path.write_bytes(head + b"\x00" * 32)
    return path


def _make_structured_docx(path: Path) -> Path:
    """LibreOffice 변환물과 동형인 '구조 있는' .docx 를 python-docx 로 만든다.

    문서 순서:
      [Heading 1] 1장 보안          → SECTION
      [Normal]    본문 문단 A        → PARAGRAPH (heading_path=("1장 보안",))
      [Heading 2] 1.1 권한          → SUBHEADING (("1장 보안","1.1 권한"))
      [Normal]    본문 문단 B        → PARAGRAPH (("1장 보안","1.1 권한"))
      [표]        2x2               → TABLE     (같은 heading_path)
      [Heading 1] 2장 배포          → SECTION   (heading_path 재설정)
      [Normal]    본문 문단 C        → PARAGRAPH (("2장 배포",))

    왜 add_heading 의 level 인자를 쓰는가: python-docx 의 add_heading(text, level=N)
    은 단락 스타일을 표준 'Heading N' 으로 지정한다 — 파서의 _heading_level 이
    인식하는 바로 그 스타일이다. 표 사이에 문단이 끼도록 본문 순서대로 추가해
    '문서 순서 보존'을 실제로 검증할 수 있게 한다."""
    doc = Document()
    doc.add_heading("1장 보안", level=1)
    doc.add_paragraph("본문 문단 A")
    doc.add_heading("1.1 권한", level=2)
    doc.add_paragraph("본문 문단 B")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "헤더1"
    table.cell(0, 1).text = "헤더2"
    table.cell(1, 0).text = "값1"
    table.cell(1, 1).text = "값2"
    doc.add_heading("2장 배포", level=1)
    doc.add_paragraph("본문 문단 C")
    doc.save(str(path))
    return path


def _primed_parser(
    soffice_cmd: str = "soffice", timeout_sec: float = 120.0
) -> HwpViaLibreOfficeParser:
    """_ensure_config 를 단락시키기 위해 _cfg 를 미리 채운 파서.

    실 HwpConfig 로딩/환경변수 의존을 피해 변환·파싱 테스트를 결정론적으로 만든다."""
    parser = HwpViaLibreOfficeParser()
    parser._cfg = (soffice_cmd, timeout_sec)
    return parser


# ─────────────────────────────────────────────
# 1) can_parse — fail-closed
# ─────────────────────────────────────────────
def test_can_parse_valid_ole_magic_returns_true(tmp_path: Path) -> None:
    """확장자 .hwp + OLE2 매직바이트(D0CF11E0A1B11AE1) 일치 → True."""
    path = _write_hwp_magic(tmp_path / "doc.hwp", valid=True)
    assert HwpViaLibreOfficeParser().can_parse(path) is True


def test_can_parse_wrong_magic_returns_false(tmp_path: Path) -> None:
    """확장자는 .hwp 지만 선두가 OLE2 시그니처가 아니면 False(fail-closed)."""
    path = _write_hwp_magic(tmp_path / "doc.hwp", valid=False)
    assert HwpViaLibreOfficeParser().can_parse(path) is False


def test_can_parse_hwpx_zip_returns_false(tmp_path: Path) -> None:
    """HWPX(개방형 ZIP 'PK\\x03\\x04')는 이 파서 소관이 아니므로 False.

    설령 확장자를 .hwp 로 위장해도 매직이 OLE2 가 아니라 ZIP 이므로 거부된다."""
    path = tmp_path / "doc.hwp"
    path.write_bytes(b"PK\x03\x04" + b"\x00" * 32)  # ZIP 시그니처
    assert HwpViaLibreOfficeParser().can_parse(path) is False


def test_can_parse_wrong_extension_returns_false(tmp_path: Path) -> None:
    """매직바이트가 OLE2 라도 확장자가 .hwp 가 아니면 False(확장자 우선 거부)."""
    path = tmp_path / "doc.hwpx"  # .hwpx — 다른 파서 소관
    path.write_bytes(_HWP_OLE_MAGIC + b"\x00" * 32)
    assert HwpViaLibreOfficeParser().can_parse(path) is False


def test_can_parse_missing_file_returns_false(tmp_path: Path) -> None:
    """없는 파일은 OSError 를 삼키고 False(fail-closed)."""
    assert HwpViaLibreOfficeParser().can_parse(tmp_path / "nope.hwp") is False


def test_supported_extensions_and_no_gpu() -> None:
    """정체성 플래그 — .hwp 만 지원, GPU 불필요(LibreOffice 변환은 CPU)."""
    parser = HwpViaLibreOfficeParser()
    assert parser.supported_extensions == (".hwp",)
    assert parser.requires_gpu is False


# ─────────────────────────────────────────────
# 2) 변환 후 파싱 — 실 .docx fixture(구조 보존)
# ─────────────────────────────────────────────
def test_parse_docx_structured_builds_nodes_in_document_order(tmp_path: Path) -> None:
    """실 .docx → SECTION/SUBHEADING/PARAGRAPH/TABLE 노드를 문서 순서대로 생성.

    핵심: doc.paragraphs/doc.tables 를 따로 도는 방식이 아니라 body 자식을
    순서대로 훑어 '표가 문단들 사이에 끼어 있던 흐름'을 그대로 복원해야 한다."""
    docx_path = _make_structured_docx(tmp_path / "converted.docx")
    parser = HwpViaLibreOfficeParser()
    warnings: list[str] = []

    nodes = parser._parse_docx_structured(docx_path, warnings)

    # 빈 문단은 생략되므로 노드 수는 정확히 7개(헤딩3 + 문단3 + 표1).
    assert warnings == []
    assert len(nodes) == 7

    types = [n.element_type for n in nodes]
    assert types == [
        ElementType.SECTION,  # 1장 보안
        ElementType.PARAGRAPH,  # 본문 문단 A
        ElementType.SUBHEADING,  # 1.1 권한
        ElementType.PARAGRAPH,  # 본문 문단 B
        ElementType.TABLE,  # 2x2 표
        ElementType.SECTION,  # 2장 배포
        ElementType.PARAGRAPH,  # 본문 문단 C
    ]

    # order 는 문서 전체 읽기 순서대로 0..6 단조 증가(빈 문단 없으므로 연속).
    assert [n.order for n in nodes] == [0, 1, 2, 3, 4, 5, 6]


def test_parse_docx_structured_accumulates_heading_path(tmp_path: Path) -> None:
    """heading_path 누적 — 본문/표가 현재 제목 경로를 맥락으로 받고, 새 SECTION 에서 재설정."""
    docx_path = _make_structured_docx(tmp_path / "converted.docx")
    parser = HwpViaLibreOfficeParser()
    nodes = parser._parse_docx_structured(docx_path, [])

    by_text = {n.text: n for n in nodes}

    # SECTION 은 자기 제목만으로 경로 재설정.
    assert by_text["1장 보안"].heading_path == ("1장 보안",)
    # SECTION 직후 본문은 그 섹션 경로를 물려받는다.
    assert by_text["본문 문단 A"].heading_path == ("1장 보안",)
    # Heading 2 는 (level-1=1) 깊이로 자른 뒤 자신을 덧붙인다.
    assert by_text["1.1 권한"].heading_path == ("1장 보안", "1.1 권한")
    # 그 아래 본문/표는 소제목 경로까지 물려받는다.
    assert by_text["본문 문단 B"].heading_path == ("1장 보안", "1.1 권한")
    # 두 번째 SECTION 은 경로를 자기 제목만으로 재설정(이전 소제목 흔적 제거).
    assert by_text["2장 배포"].heading_path == ("2장 배포",)
    assert by_text["본문 문단 C"].heading_path == ("2장 배포",)


def test_parse_docx_structured_table_serialized_with_pipe(tmp_path: Path) -> None:
    """표 노드는 셀을 ' | ' 로, 행을 줄바꿈으로 직렬화한다(다른 파서와 동일 포맷)."""
    docx_path = _make_structured_docx(tmp_path / "converted.docx")
    parser = HwpViaLibreOfficeParser()
    nodes = parser._parse_docx_structured(docx_path, [])

    table_nodes = [n for n in nodes if n.element_type == ElementType.TABLE]
    assert len(table_nodes) == 1
    assert table_nodes[0].text == "헤더1 | 헤더2\n값1 | 값2"
    # 표는 직전까지의 heading_path(소제목)를 맥락으로 받는다.
    assert table_nodes[0].heading_path == ("1장 보안", "1.1 권한")


def test_parse_docx_structured_skips_empty_paragraphs(tmp_path: Path) -> None:
    """공백뿐인 문단은 노드를 만들지 않는다(검색 노이즈 방지)."""
    doc = Document()
    doc.add_paragraph("실제 내용")
    doc.add_paragraph("   ")  # 공백만 — 생략 대상
    doc.add_paragraph("")  # 빈 문단 — 생략 대상
    docx_path = tmp_path / "sparse.docx"
    doc.save(str(docx_path))

    parser = HwpViaLibreOfficeParser()
    nodes = parser._parse_docx_structured(docx_path, [])

    assert len(nodes) == 1
    assert nodes[0].text == "실제 내용"


def test_parse_docx_structured_corrupted_docx_fails_soft(tmp_path: Path) -> None:
    """변환물(.docx)이 손상됐으면 빈 목록 + 경고로 환원한다(예외 비전파)."""
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"PK\x03\x04 not-a-real-docx-zip")
    parser = HwpViaLibreOfficeParser()
    warnings: list[str] = []

    nodes = parser._parse_docx_structured(broken, warnings)

    assert nodes == []
    assert any("docx 열기 실패" in w for w in warnings)


@pytest.mark.asyncio
async def test_parse_korean_heading_styles(tmp_path: Path) -> None:
    """한글 '제목 N' 스타일도 SECTION/SUBHEADING 으로 인식한다(가능 범위).

    python-docx 로 표준 'Heading N' 스타일이 부여된 단락의 style.name 을 한글
    'Title/제목 N' 으로 직접 바꿔, _heading_level 의 한글 prefix 인식 경로를 태운다.
    (LibreOffice 한국어 로캘 변환물에서 한글 스타일명이 올 수 있음을 모사.)"""
    doc = Document()
    h1 = doc.add_heading("대제목", level=1)
    h2 = doc.add_heading("소제목", level=2)
    docx_path = tmp_path / "kor.docx"
    # 스타일 객체의 이름을 한글로 바꾼다(파서는 style.name 문자열만 본다).
    h1.style.name = "제목 1"
    h2.style.name = "제목 2"
    doc.save(str(docx_path))

    parser = HwpViaLibreOfficeParser()
    nodes = parser._parse_docx_structured(docx_path, [])

    types = [n.element_type for n in nodes]
    assert ElementType.SECTION in types
    assert ElementType.SUBHEADING in types


# ─────────────────────────────────────────────
# 3) _heading_level — 스타일명 → 레벨
# ─────────────────────────────────────────────
@pytest.mark.parametrize(
    ("style_name", "expected"),
    [
        ("Heading 1", 1),
        ("Heading 2", 2),
        ("Heading 3", 3),
        ("heading 1", 1),  # 대소문자 무관
        ("제목 1", 1),  # 한글 prefix
        ("제목 2", 2),
        ("Normal", 0),  # 본문
        ("Title", 0),  # 표지 제목은 본문(0) — 과탐 방지
        ("", 0),  # 무명 스타일
        ("List Paragraph", 0),  # 숫자 없는 일반 스타일
    ],
)
def test_heading_level_classification(style_name: str, expected: int) -> None:
    """스타일 이름에서 제목 레벨(0=본문, N=제목 깊이)을 추출한다."""
    assert HwpViaLibreOfficeParser._heading_level(style_name) == expected


def test_heading_level_clamps_deep_levels() -> None:
    """비정상적으로 깊은 레벨(Heading 99)은 9 로 클램프된다(heading_path 폭주 방지)."""
    assert HwpViaLibreOfficeParser._heading_level("Heading 99") == 9
    assert HwpViaLibreOfficeParser._heading_level("Heading 9") == 9
    assert HwpViaLibreOfficeParser._heading_level("Heading 10") == 9


# ─────────────────────────────────────────────
# 4) _table_node — 직렬화/빈 표/셀 실패 흡수
# ─────────────────────────────────────────────
def test_table_node_serializes_rows_and_cells(tmp_path: Path) -> None:
    """정상 표 → ' | ' 로 셀을, 줄바꿈으로 행을 이어 직렬화한 TABLE 노드."""
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    cells = [["a", "b", "c"], ["d", "e", "f"]]
    for r in range(2):
        for c in range(3):
            table.cell(r, c).text = cells[r][c]
    docx_path = tmp_path / "t.docx"
    doc.save(str(docx_path))
    # 저장/재로드해 실제 docx 표 객체로 검증.
    table_obj = Document(str(docx_path)).tables[0]

    node = HwpViaLibreOfficeParser._table_node(table_obj, order=5, heading_path=("X",), warnings=[])

    assert node is not None
    assert node.element_type == ElementType.TABLE
    assert node.text == "a | b | c\nd | e | f"
    assert node.order == 5
    assert node.heading_path == ("X",)


def test_table_node_empty_table_returns_none(tmp_path: Path) -> None:
    """모든 셀이 공백인 표는 None(노드 생략 — 빈 격자 노이즈 방지)."""
    doc = Document()
    doc.add_table(rows=2, cols=2)  # 모든 셀 공백
    docx_path = tmp_path / "empty.docx"
    doc.save(str(docx_path))
    table_obj = Document(str(docx_path)).tables[0]

    node = HwpViaLibreOfficeParser._table_node(table_obj, order=0, heading_path=(), warnings=[])
    assert node is None


def test_table_node_row_access_failure_returns_none() -> None:
    """표 행(rows) 접근이 깨지면 경고 후 None(표 전체 포기, 예외 비전파).

    rows 속성 접근에서 AttributeError 를 던지는 가짜 표로 fail-soft 를 검증한다."""

    class _BrokenTable:
        @property
        def rows(self):  # noqa: ANN202 — 테스트용 가짜
            raise AttributeError("rows 접근 불가")

    warnings: list[str] = []
    node = HwpViaLibreOfficeParser._table_node(
        _BrokenTable(), order=0, heading_path=(), warnings=warnings
    )
    assert node is None
    assert any("표 행 접근 실패" in w for w in warnings)


def test_table_node_cell_access_failure_absorbed_as_blank() -> None:
    """셀 1개 접근 실패는 빈칸으로 흡수하고 나머지 셀은 직렬화한다(표 전체를 버리지 않음).

    한 셀의 .text 접근이 ValueError 를 던지는 가짜 표 구조로, 그 셀만 ''(빈칸)으로
    처리되고 나머지 셀은 보존되는지 확인한다."""

    class _Cell:
        def __init__(self, value: str | None, *, broken: bool = False) -> None:
            self._value = value
            self._broken = broken

        @property
        def text(self) -> str:
            if self._broken:
                raise ValueError("셀 텍스트 접근 불가")
            return self._value or ""

    class _Row:
        def __init__(self, cells: list[_Cell]) -> None:
            self.cells = cells

    class _Table:
        def __init__(self, rows: list[_Row]) -> None:
            self.rows = rows

    fake = _Table([_Row([_Cell("ok"), _Cell(None, broken=True), _Cell("tail")])])

    node = HwpViaLibreOfficeParser._table_node(fake, order=0, heading_path=(), warnings=[])
    assert node is not None
    # 깨진 가운데 셀은 빈칸으로 흡수 → "ok |  | tail".
    assert node.text == "ok |  | tail"


# ─────────────────────────────────────────────
# 5) fail-soft — _convert_to_docx subprocess 예외 주입
# ─────────────────────────────────────────────
@pytest.mark.asyncio
async def test_parse_soffice_not_installed_fails_soft(tmp_path: Path) -> None:
    """soffice 미설치(FileNotFoundError) → 빈 트리 + 경고(예외 비전파)."""
    hwp = _write_hwp_magic(tmp_path / "doc.hwp")
    parser = _primed_parser(soffice_cmd="no-such-soffice")

    with patch("subprocess.run", side_effect=FileNotFoundError("soffice 없음")):
        tree = await parser.parse(hwp)

    assert tree.doc_format == "hwp"
    assert tree.title == "doc"
    assert tree.nodes == ()
    assert any("실행 파일을 찾지 못함" in w for w in tree.warnings)


@pytest.mark.asyncio
async def test_parse_soffice_convert_error_fails_soft(tmp_path: Path) -> None:
    """변환 비정상 종료(CalledProcessError) → 빈 트리 + rc/stderr 경고."""
    hwp = _write_hwp_magic(tmp_path / "doc.hwp")
    parser = _primed_parser()

    err = subprocess.CalledProcessError(returncode=1, cmd=["soffice"], stderr=b"conversion boom")
    with patch("subprocess.run", side_effect=err):
        tree = await parser.parse(hwp)

    assert tree.nodes == ()
    assert any("변환 실패" in w for w in tree.warnings)


@pytest.mark.asyncio
async def test_parse_soffice_timeout_fails_soft(tmp_path: Path) -> None:
    """변환 시간 초과(TimeoutExpired) → 빈 트리 + 타임아웃 경고."""
    hwp = _write_hwp_magic(tmp_path / "doc.hwp")
    parser = _primed_parser(timeout_sec=30.0)

    timeout_err = subprocess.TimeoutExpired(cmd="soffice", timeout=30.0)
    with patch("subprocess.run", side_effect=timeout_err):
        tree = await parser.parse(hwp)

    assert tree.nodes == ()
    assert any("시간 초과" in w for w in tree.warnings)


@pytest.mark.asyncio
async def test_parse_converted_docx_not_found_fails_soft(tmp_path: Path) -> None:
    """soffice 가 0 으로 정상 종료했지만 변환물(.docx)이 outdir 에 없으면 경고 + 빈 트리.

    subprocess.run 은 성공(반환값)으로 mock 하되 실제 .docx 는 생성하지 않으므로,
    파서의 '변환물 미발견' 분기를 태운다."""
    hwp = _write_hwp_magic(tmp_path / "doc.hwp")
    parser = _primed_parser()

    ok = SimpleNamespace(returncode=0, stdout=b"converted", stderr=b"")
    with patch("subprocess.run", return_value=ok):
        tree = await parser.parse(hwp)

    assert tree.nodes == ()
    assert any("찾지 못함" in w for w in tree.warnings)


@pytest.mark.asyncio
async def test_parse_full_path_with_produced_docx(tmp_path: Path) -> None:
    """변환 성공 모사 — subprocess.run 이 outdir 에 실 .docx 를 떨군 것처럼 동작시켜
    parse() 전 구간(변환→파싱→노드 트리)을 mock 경계에서 e2e 로 검증한다.

    실 soffice/실 .hwp 없이도, '변환물이 생긴 뒤'부터의 파싱 경로가 정상 트리를
    만드는지 확인한다(변환 호출만 mock, 파싱은 실 python-docx)."""
    hwp = _write_hwp_magic(tmp_path / "report.hwp")
    parser = _primed_parser()

    def _fake_run(cmd, **kwargs):  # noqa: ANN001, ANN202 — 테스트용 fake
        # cmd = [soffice, --headless, --convert-to, docx, --outdir, <outdir>, <hwp>]
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        # LibreOffice 가 입력 stem 으로 .docx 를 떨구는 것을 모사.
        _make_structured_docx(outdir / "report.docx")
        return SimpleNamespace(returncode=0, stdout=b"ok", stderr=b"")

    with patch("subprocess.run", side_effect=_fake_run):
        tree = await parser.parse(hwp)

    assert tree.doc_format == "hwp"
    assert tree.title == "report"
    assert tree.warnings == ()
    types = [n.element_type for n in tree.nodes]
    assert ElementType.SECTION in types
    assert ElementType.SUBHEADING in types
    assert ElementType.TABLE in types
    assert ElementType.PARAGRAPH in types


# ─────────────────────────────────────────────
# 6) config / 환경변수 / timeout 보정
# ─────────────────────────────────────────────
def test_hwp_config_defaults() -> None:
    """HwpConfig 기본값 — soffice_cmd 는 Windows 설치 경로, timeout 120 초."""
    cfg = HwpConfig()
    assert cfg.soffice_cmd == _DEFAULT_SOFFICE_CMD
    assert cfg.soffice_cmd.endswith("soffice.exe")
    assert cfg.convert_timeout_sec == 120.0


def test_ensure_config_reads_defaults() -> None:
    """_ensure_config 가 HwpConfig 기본값을 읽어 (soffice_cmd, timeout) 으로 캐시한다."""
    parser = HwpViaLibreOfficeParser()
    soffice_cmd, timeout_sec = parser._ensure_config()
    assert soffice_cmd.endswith("soffice.exe")
    assert timeout_sec == 120.0


def test_ensure_config_caches_result() -> None:
    """_ensure_config 는 첫 호출 결과를 캐시해 동일 튜플을 재반환한다(지연 1회)."""
    parser = HwpViaLibreOfficeParser()
    first = parser._ensure_config()
    second = parser._ensure_config()
    assert first is second


def test_ensure_config_env_soffice_fallback(monkeypatch: pytest.MonkeyPatch) -> None:
    """config.soffice_cmd 가 비면 NEXUS_SOFFICE_CMD 환경변수로 폴백한다.

    실 config 의 기본값은 비어 있지 않으므로, load_and_validate_config 를 patch 해
    soffice_cmd='' 인 config 를 돌려주게 만들어 '빈 config → 환경변수' 폴백 경로를
    결정론적으로 태운다(빈 문자열 환경변수의 OS/세팅 의존성을 회피)."""
    monkeypatch.setenv("NEXUS_SOFFICE_CMD", r"C:\fallback\soffice.exe")
    fake_cfg = SimpleNamespace(hwp=SimpleNamespace(soffice_cmd="", convert_timeout_sec=120.0))

    parser = HwpViaLibreOfficeParser()
    with patch("core.config.load_and_validate_config", return_value=fake_cfg):
        soffice_cmd, _timeout = parser._ensure_config()

    assert soffice_cmd == r"C:\fallback\soffice.exe"


def test_ensure_config_default_when_all_empty(monkeypatch: pytest.MonkeyPatch) -> None:
    """config.soffice_cmd 와 NEXUS_SOFFICE_CMD 가 모두 비면 개발 기본값으로 최종 폴백."""
    monkeypatch.delenv("NEXUS_SOFFICE_CMD", raising=False)
    fake_cfg = SimpleNamespace(hwp=SimpleNamespace(soffice_cmd="", convert_timeout_sec=120.0))

    parser = HwpViaLibreOfficeParser()
    with patch("core.config.load_and_validate_config", return_value=fake_cfg):
        soffice_cmd, _timeout = parser._ensure_config()

    assert soffice_cmd == _DEFAULT_SOFFICE_CMD


def test_ensure_config_invalid_timeout_falls_back_to_default(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """timeout 이 0/음수면 기본값(_DEFAULT_TIMEOUT_SEC)으로 보정한다(하한 보호)."""
    fake_cfg = SimpleNamespace(
        hwp=SimpleNamespace(soffice_cmd=r"C:\x\soffice.exe", convert_timeout_sec=0.0)
    )

    parser = HwpViaLibreOfficeParser()
    with patch("core.config.load_and_validate_config", return_value=fake_cfg):
        _soffice_cmd, timeout_sec = parser._ensure_config()

    assert timeout_sec == _DEFAULT_TIMEOUT_SEC


def test_ensure_config_load_failure_falls_back_to_env(monkeypatch: pytest.MonkeyPatch) -> None:
    """HwpConfig 로딩이 예외로 실패해도 환경변수 폴백으로 soffice 경로를 확보한다(부분 가용성)."""
    monkeypatch.setenv("NEXUS_SOFFICE_CMD", r"C:\env\soffice.exe")

    parser = HwpViaLibreOfficeParser()
    with patch(
        "core.config.load_and_validate_config",
        side_effect=RuntimeError("config 시스템 고장"),
    ):
        soffice_cmd, timeout_sec = parser._ensure_config()

    assert soffice_cmd == r"C:\env\soffice.exe"
    assert timeout_sec == _DEFAULT_TIMEOUT_SEC
